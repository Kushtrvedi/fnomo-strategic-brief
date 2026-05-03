#!/usr/bin/env python3
"""Record Paperclip orchestration setup in the Fnomo workbook.

Sheet remains source of truth. Paperclip is recorded as orchestration/visibility
in the task queue and execution history; it is not treated as CRM authority.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from shutil import copy2
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "EXCEL" / "War Room_ Community Outreach Pipeline - FNOMO Master.xlsx"


def compact(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def find_header_row(ws: Any, required: str) -> tuple[int, dict[str, int]]:
    for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = [compact(cell) for cell in row]
        if required in values:
            return row_index, {header: idx + 1 for idx, header in enumerate(values) if header}
    raise RuntimeError(f"Missing header {required}")


def upsert_task(ws: Any, headers: dict[str, int], task_id: str, values: dict[str, str]) -> None:
    row_index = None
    for idx in range(1, ws.max_row + 1):
        if compact(ws.cell(idx, headers["Task_ID"]).value) == task_id:
            row_index = idx
            break
    row_index = row_index or ws.max_row + 1
    values = {"Task_ID": task_id, **values}
    for header, value in values.items():
        if header in headers:
            ws.cell(row_index, headers[header]).value = value


def append_history(wb: Any, operating_date: date) -> None:
    ws = wb["FNOMO_EXECUTION_HISTORY"]
    marker = "Paperclip Orchestration Setup"
    for row_index in range(ws.max_row, 1, -1):
        if compact(ws.cell(row_index, 1).value) == operating_date.isoformat() and compact(ws.cell(row_index, 2).value) == marker:
            ws.delete_rows(row_index, 1)
    row_index = ws.max_row + 1
    values = [
        operating_date.isoformat(),
        marker,
        "Fnomo Paperclip Operating System",
        "Codex",
        "Paperclip + Workbook",
        "Installed/verified Paperclip local dashboard, imported PA plus specialist agents, created task types, added system-health monitor, registered Windows Startup auto-start, and recorded Monday follow-up guardrails.",
        "Use Paperclip for orchestration only; verify auto-start after Codex/system restart; execute Monday Follow-up 1 from the workbook source of truth.",
        "Main chat PA control",
    ]
    for col_index, value in enumerate(values, start=1):
        ws.cell(row_index, col_index).value = value


def main() -> None:
    operating_date = date.fromisoformat("2026-05-03")
    monday = date.fromisoformat("2026-05-04")
    backup = WORKBOOK.parent / "backups" / f"{WORKBOOK.stem}.before-paperclip-setup.{datetime.now().strftime('%Y%m%d-%H%M%S')}{WORKBOOK.suffix}"
    backup.parent.mkdir(exist_ok=True)
    copy2(WORKBOOK, backup)

    wb = load_workbook(WORKBOOK)
    tasks = wb["FNOMO_EXECUTION_TASKS"]
    _, headers = find_header_row(tasks, "Task_ID")
    upsert_task(
        tasks,
        headers,
        "TDA-PAPERCLIP-001",
        {
            "Task": "Paperclip Fnomo OS Setup",
            "Owner": "Codex",
            "Priority": "A",
            "Next_Action": "Use Paperclip as orchestration layer only; Sheet remains source of truth.",
            "Deadline": operating_date.isoformat(),
            "Status": "Done",
            "Related_Segment": "Paperclip",
            "Notes": "Local dashboard verified at http://127.0.0.1:3100; Fnomo specialist agents and task types imported.",
        },
    )
    upsert_task(
        tasks,
        headers,
        "TDA-PAPERCLIP-HEALTH-001",
        {
            "Task": "Paperclip System Health Monitor",
            "Owner": "Codex / System Health Monitor",
            "Priority": "A",
            "Next_Action": "Run start-of-day, mid-cycle, and end-of-day checks; unblock stale tasks and enforce SLA.",
            "Deadline": monday.isoformat(),
            "Status": "Active",
            "Related_Segment": "Paperclip",
            "Notes": "System Check output required after major cycles: status, issue, correction applied, next priority.",
        },
    )
    upsert_task(
        tasks,
        headers,
        "TDA-PAPERCLIP-FU1-001",
        {
            "Task": "Paperclip Monday Follow-up 1 Queue",
            "Owner": "Follow-up Agent / Codex",
            "Priority": "A",
            "Next_Action": "Execute 22 Tier A BO/CA Messaging Test Batch Follow-up 1 messages on Monday only.",
            "Deadline": monday.isoformat(),
            "Status": "Planned for Monday",
            "Related_Segment": "Messaging Test Batch 1",
            "Notes": "Variant split: Process Audit 15, Advisor/CA 5, Operator/BO 2. Platform rows held for separate angle.",
        },
    )
    upsert_task(
        tasks,
        headers,
        "TDA-PAPERCLIP-AUTOSTART-001",
        {
            "Task": "Paperclip Windows Auto Startup",
            "Owner": "Codex",
            "Priority": "A",
            "Next_Action": "Verify Windows Startup shortcut starts Paperclip at user logon and health endpoint returns ok after restart.",
            "Deadline": operating_date.isoformat(),
            "Status": "Done",
            "Related_Segment": "Paperclip",
            "Notes": "Created user Startup shortcut Fnomo Paperclip AutoStart.lnk to run scripts/start_paperclip_autostart.ps1 at logon. Task Scheduler creation was blocked by Windows permissions.",
        },
    )
    append_history(wb, operating_date)
    wb.save(WORKBOOK)
    print(
        {
            "backup": str(backup),
            "tasks_upserted": [
                "TDA-PAPERCLIP-001",
                "TDA-PAPERCLIP-HEALTH-001",
                "TDA-PAPERCLIP-FU1-001",
                "TDA-PAPERCLIP-AUTOSTART-001",
            ],
            "history_logged": True,
        }
    )


if __name__ == "__main__":
    main()
