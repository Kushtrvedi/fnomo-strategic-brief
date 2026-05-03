#!/usr/bin/env python3
"""
Prepare Messaging Test Batch Follow-up 1 variants for Monday execution.

Sunday rule: planning only. This script updates workbook/Notion planning
state, message bank, tasks, and history. It does not send outreach.
"""

from __future__ import annotations

import json
import os
import re
import sys
import time
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from shutil import copy2
from typing import Any

import requests
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "EXCEL" / "War Room_ Community Outreach Pipeline - FNOMO Master.xlsx"
CLIENTS_DB = "e8a8aa3c-b995-8368-8495-014b8ea295c3"
CONTACTS_DB = "3548aa3c-b995-8148-a653-fead444e56fe"
EXECUTION_DB = "6fb8aa3c-b995-82ec-806b-01b6cd1b2915"
HISTORY_DB = "3548aa3c-b995-818a-8ede-ed9a5d99aac3"
NOTION_VERSION = "2022-06-28"
SESSION = requests.Session()
SCHEMA_CACHE: dict[str, set[str]] = {}

OPERATING_TAG = "Messaging Test Batch 1"
FOLLOWUP_TAG = "Follow-up 1"
PLANNING_TAG = "Monday Planning"
PLATFORM_HOLD_TAG = "Platform Follow-up Separate Angle"
PLATFORM_TERMS = {
    "blinkit",
    "practo",
    "swiggy",
    "zomato",
    "nobroker",
    "internshala",
    "collegedunia",
    "housing",
    "mygate",
}


VARIANTS: dict[str, dict[str, str]] = {
    "Follow-up 1 - Loss Recall": {
        "channel": "WhatsApp / LinkedIn DM",
        "subject": "Quick follow-up",
        "message": (
            "Quick follow-up -\n\n"
            "In the last few months, was there a decision that looked right at the time but you later questioned?\n\n"
            "Not about finding ideas - just the moment before acting.\n\n"
            "Curious how you handle that step today."
        ),
        "rule": "Default when profile is broad Tier A and no sharper role cue is available.",
    },
    "Follow-up 1 - Process Audit": {
        "channel": "WhatsApp / LinkedIn DM",
        "subject": "Adding context to my last note",
        "message": (
            "Adding context to my last note -\n\n"
            "Most people I speak to have strong research, but no structured step to validate the decision before acting.\n\n"
            "In your process, does that step exist, or does it move straight from analysis to action?"
        ),
        "rule": "Use for analytical, technical, listed-company, industrial, finance-heavy, or process-led profiles.",
    },
    "Follow-up 1 - Advisor CA": {
        "channel": "WhatsApp / LinkedIn DM",
        "subject": "Following up",
        "message": (
            "Following up -\n\n"
            "Have you had a client recently take a decision that didn't work out as expected?\n\n"
            "Usually the gap isn't advice - it's what happens just before they act.\n\n"
            "Do you have a validation step there, or is it still a gap?"
        ),
        "rule": "Use for CA, advisor, branch, association, and professional trust routes.",
    },
    "Follow-up 1 - Operator BO": {
        "channel": "WhatsApp / LinkedIn DM",
        "subject": "Quick one",
        "message": (
            "Quick one -\n\n"
            "Before deploying capital, do you run a structured check on the decision itself, or rely on inputs + gut?\n\n"
            "Most mistakes I see aren't bad information - it's unvalidated action."
        ),
        "rule": "Use for owner-led business, operator, HNI, and capital deployment routes.",
    },
    "Follow-up 1 - Close The Loop": {
        "channel": "WhatsApp / LinkedIn DM",
        "subject": "Closing the loop",
        "message": (
            "Closing the loop on my earlier messages.\n\n"
            "If this isn't relevant right now, all good.\n\n"
            "If the \"validate before acting\" step is something you've seen missing, happy to compare notes briefly."
        ),
        "rule": "Use on Day 7 if still silent. Do not use for Monday 48h follow-up unless a lead is already 7 days silent.",
    },
}


def die(message: str) -> None:
    print(f"ERROR: {message}", file=sys.stderr)
    raise SystemExit(1)


def compact(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    return str(value).strip()


def norm(value: Any) -> str:
    return " ".join(compact(value).casefold().split())


def parse_date(value: Any) -> str:
    text = compact(value)
    if not text:
        return ""
    if len(text) >= 10 and text[4] == "-" and text[7] == "-":
        return text[:10]
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            pass
    return text


def headers() -> dict[str, str]:
    token = os.environ.get("NOTION_API_KEY")
    if not token:
        die("NOTION_API_KEY missing")
    return {
        "Authorization": f"Bearer {token}",
        "Notion-Version": NOTION_VERSION,
        "Content-Type": "application/json",
    }


def notion(method: str, path: str, payload: dict[str, Any] | None = None) -> dict[str, Any]:
    for attempt in range(6):
        response = SESSION.request(
            method,
            f"https://api.notion.com/v1/{path.lstrip('/')}",
            headers=headers(),
            json=payload,
            timeout=60,
        )
        if response.status_code == 429:
            time.sleep(2 + attempt)
            continue
        if response.status_code >= 400:
            raise RuntimeError(f"{method} {path} failed {response.status_code}: {response.text[:1500]}")
        return response.json()
    raise RuntimeError(f"{method} {path} failed after retries")


def query_database(database_id: str, filter_payload: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    cursor = None
    while True:
        payload: dict[str, Any] = {"page_size": 100}
        if cursor:
            payload["start_cursor"] = cursor
        if filter_payload:
            payload["filter"] = filter_payload
        data = notion("POST", f"databases/{database_id}/query", payload)
        rows.extend(data.get("results", []))
        if not data.get("has_more"):
            return rows
        cursor = data.get("next_cursor")


def plain(prop: dict[str, Any] | None) -> str:
    if not prop:
        return ""
    prop_type = prop.get("type")
    if prop_type in {"title", "rich_text"}:
        return "".join(part.get("plain_text", "") for part in prop.get(prop_type, [])).strip()
    if prop_type == "select":
        return ((prop.get("select") or {}).get("name") or "").strip()
    if prop_type == "status":
        return ((prop.get("status") or {}).get("name") or "").strip()
    if prop_type == "date" and prop.get("date"):
        return (prop["date"].get("start") or "").strip()
    if prop_type == "multi_select":
        return "; ".join(item.get("name", "") for item in prop.get("multi_select", []))
    if prop_type == "relation":
        return ",".join(item.get("id", "") for item in prop.get("relation", []) if item.get("id"))
    return ""


def prop(page: dict[str, Any], name: str) -> dict[str, Any] | None:
    return (page.get("properties") or {}).get(name)


def title_of(page: dict[str, Any]) -> str:
    for value in (page.get("properties") or {}).values():
        if value.get("type") == "title":
            return plain(value)
    return ""


def relation_ids(page: dict[str, Any], name: str) -> set[str]:
    return {item.get("id", "") for item in ((prop(page, name) or {}).get("relation") or []) if item.get("id")}


def title_text(value: str) -> dict[str, Any]:
    return {"title": [{"text": {"content": value[:1800]}}]}


def rich_text(value: Any) -> dict[str, Any]:
    text = compact(value)[:1900]
    return {"rich_text": [{"text": {"content": text}}]} if text else {"rich_text": []}


def date_value(value: Any) -> dict[str, Any]:
    text = compact(value)
    return {"date": {"start": text}} if text else {"date": None}


def select_value(value: Any) -> dict[str, Any]:
    text = compact(value)[:100]
    return {"select": {"name": text}} if text else {"select": None}


def status_value(value: Any) -> dict[str, Any]:
    text = compact(value)[:100]
    return {"status": {"name": text}} if text else {"status": None}


def multi_select(existing_prop: dict[str, Any] | None, tags: list[str]) -> dict[str, Any]:
    existing = []
    if existing_prop and existing_prop.get("type") == "multi_select":
        existing = [item.get("name", "") for item in existing_prop.get("multi_select", []) if item.get("name")]
    for tag in tags:
        if tag and tag not in existing:
            existing.append(tag)
    return {"multi_select": [{"name": tag} for tag in existing]}


def allowed_props(database_id: str, props: dict[str, Any]) -> dict[str, Any]:
    if database_id not in SCHEMA_CACHE:
        SCHEMA_CACHE[database_id] = set((notion("GET", f"databases/{database_id}").get("properties") or {}).keys())
    return {key: value for key, value in props.items() if key in SCHEMA_CACHE[database_id]}


def append_note(existing: Any, note: str) -> str:
    text = compact(existing)
    if note in text:
        return text
    return (text + " | " + note).strip(" |")


def strip_stale_platform_variant_note(existing: Any) -> str:
    text = compact(existing)
    return re.sub(
        r"\s*\|\s*\[2026-05-03 Sunday planning\] Follow-up 1 variant locked for 2026-05-04; no links, no product, no pitch\. Variant: Follow-up 1 - [^.]+\.",
        "",
        text,
    ).strip(" |")


def find_header_row(ws: Any, required: str) -> tuple[int, dict[str, int]]:
    for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
        headers = [compact(cell) for cell in row]
        if required in headers:
            return row_index, {header: idx + 1 for idx, header in enumerate(headers) if header}
    die(f"Missing header {required}")


def infer_variant(lead_type: str, notes: str, company: str) -> str:
    blob = " ".join([lead_type, notes, company]).lower()
    analytical_terms = [
        "listed",
        "bse",
        "technical",
        "industrial",
        "transformer",
        "vfd",
        "scr",
        "manufactur",
        "procurement",
        "capex",
        "financial",
        "process",
    ]
    if any(term in blob for term in analytical_terms):
        return "Follow-up 1 - Process Audit"
    if "ca" in lead_type.lower() or "association" in lead_type.lower() or "advisor" in blob or "icai" in blob:
        return "Follow-up 1 - Advisor CA"
    if "business" in lead_type.lower() or "owner" in blob or "operator" in blob:
        return "Follow-up 1 - Operator BO"
    return "Follow-up 1 - Loss Recall"


def action_for_variant(variant: str) -> str:
    if variant == "Follow-up 1 - Advisor CA":
        return "Follow-up 1: send Advisor/CA variant; trigger client wrong-decision recall and ask whether a validation step exists before action."
    if variant == "Follow-up 1 - Operator BO":
        return "Follow-up 1: send Operator/Business Owner variant; ask whether the decision itself is structurally checked before capital moves."
    if variant == "Follow-up 1 - Process Audit":
        return "Follow-up 1: send Process Audit variant; ask whether the process has a validation step or moves from analysis to action."
    return "Follow-up 1: send Loss Recall variant; ask what decision looked right at the time but was later questioned."


def target_row(row: dict[str, Any]) -> bool:
    if compact(row["tier"]).upper() != "A":
        return False
    if compact(row["stage"]) != "Contacted":
        return False
    tags = compact(row["tags"])
    notes = compact(row["notes"])
    if OPERATING_TAG not in tags and OPERATING_TAG not in notes:
        return False
    lead_type = compact(row["lead_type"]).lower()
    return any(token in lead_type for token in ["ca", "business"])


def platform_hold_row(row: dict[str, Any]) -> bool:
    if compact(row["tier"]).upper() != "A":
        return False
    if compact(row["stage"]) != "Contacted":
        return False
    if OPERATING_TAG not in compact(row["tags"]) and OPERATING_TAG not in compact(row["notes"]):
        return False
    blob = " ".join([row["company"], row["notes"], row["lead_type"]]).lower()
    return any(term in blob for term in PLATFORM_TERMS)


def read_pipeline_rows(ws: Any, headers: dict[str, int]) -> list[dict[str, Any]]:
    rows = []
    for row_idx in range(1, ws.max_row + 1):
        rows.append(
            {
                "row_idx": row_idx,
                "lead_id": compact(ws.cell(row_idx, headers["Lead_ID"]).value),
                "name": compact(ws.cell(row_idx, headers["Name"]).value),
                "company": compact(ws.cell(row_idx, headers["Company"]).value),
                "lead_type": compact(ws.cell(row_idx, headers["Type (CA / Business / Association)"]).value),
                "tier": compact(ws.cell(row_idx, headers["Tier (A/B/C)"]).value),
                "stage": compact(ws.cell(row_idx, headers["Stage (New / Contacted / Engaged / Qualified / Converted / Lost)"]).value),
                "last_contact": parse_date(ws.cell(row_idx, headers["Last_Contact_Date"]).value),
                "next_action": compact(ws.cell(row_idx, headers["Next_Action"]).value),
                "next_date": parse_date(ws.cell(row_idx, headers["Next_Action_Date"]).value),
                "status": compact(ws.cell(row_idx, headers["Status (Active / Waiting / No Response / Closed)"]).value),
                "notes": compact(ws.cell(row_idx, headers["Notes (free text)"]).value),
                "tags": compact(ws.cell(row_idx, headers.get("Tags", 0)).value) if "Tags" in headers else "",
            }
        )
    return rows


def apply_workbook(operating_date: date, execution_date: date) -> dict[str, Any]:
    backup = WORKBOOK.parent / "backups" / f"{WORKBOOK.stem}.before-followup1-variants.{datetime.now().strftime('%Y%m%d-%H%M%S')}{WORKBOOK.suffix}"
    backup.parent.mkdir(exist_ok=True)
    copy2(WORKBOOK, backup)
    wb = load_workbook(WORKBOOK)

    pipeline = wb["FNOMO_MASTER_PIPELINE"]
    _, headers = find_header_row(pipeline, "Lead_ID")
    rows = read_pipeline_rows(pipeline, headers)
    touched = []
    held = []
    variant_counts: Counter = Counter()
    planning_note = (
        f"[{operating_date.isoformat()} Sunday planning] Follow-up 1 variant locked for "
        f"{execution_date.isoformat()}; no links, no product, no pitch."
    )
    for row in rows:
        if platform_hold_row(row) and not target_row(row):
            row_idx = row["row_idx"]
            action = "Hold outside BO/CA Follow-up 1 variant batch; prepare a separate platform-partnership validation angle before Monday execution."
            pipeline.cell(row_idx, headers["Next_Action"]).value = action
            pipeline.cell(row_idx, headers["Next_Action_Date"]).value = execution_date.isoformat()
            pipeline.cell(row_idx, headers["Notes (free text)"]).value = append_note(
                strip_stale_platform_variant_note(row["notes"]),
                f"[{operating_date.isoformat()} Sunday planning] Excluded from BO/CA Follow-up 1 variants; platform route needs separate angle.",
            )
            if "Tags" in headers:
                tags = [
                    tag.strip()
                    for tag in (row["tags"] + f"; {PLATFORM_HOLD_TAG}; {PLANNING_TAG}").split(";")
                    if tag.strip() and tag.strip() != FOLLOWUP_TAG
                ]
                pipeline.cell(row_idx, headers["Tags"]).value = "; ".join(dict.fromkeys(tags))
            held.append({**row, "next_action": action, "next_date": execution_date.isoformat(), "variant": "Platform hold - separate angle"})
            continue
        if not target_row(row):
            continue
        variant = infer_variant(row["lead_type"], row["notes"], row["company"])
        action = action_for_variant(variant)
        row_idx = row["row_idx"]
        pipeline.cell(row_idx, headers["Next_Action"]).value = action
        pipeline.cell(row_idx, headers["Next_Action_Date"]).value = execution_date.isoformat()
        pipeline.cell(row_idx, headers["Notes (free text)"]).value = append_note(row["notes"], f"{planning_note} Variant: {variant}.")
        if "Tags" in headers:
            tags = [tag.strip() for tag in (row["tags"] + f"; {FOLLOWUP_TAG}; {PLANNING_TAG}").split(";") if tag.strip()]
            pipeline.cell(row_idx, headers["Tags"]).value = "; ".join(dict.fromkeys(tags))
        touched.append({**row, "variant": variant, "next_action": action, "next_date": execution_date.isoformat()})
        variant_counts[variant] += 1

    update_outreach_ready(wb, operating_date)
    update_execution_task(wb, operating_date, execution_date, len(touched), variant_counts)
    append_history(wb, operating_date, execution_date, len(touched), variant_counts, len(held))
    wb.save(WORKBOOK)
    return {"backup": str(backup), "touched": touched, "held": held, "variant_counts": dict(variant_counts)}


def update_outreach_ready(wb: Any, operating_date: date) -> None:
    ws = wb["FNOMO_OUTREACH_READY"]
    _, headers = find_header_row(ws, "Lead_Name")
    existing = {compact(ws.cell(row_idx, headers["Lead_Name"]).value): row_idx for row_idx in range(1, ws.max_row + 1)}
    for name, payload in VARIANTS.items():
        row_idx = existing.get(name) or ws.max_row + 1
        values = {
            "Lead_Name": name,
            "Channel": payload["channel"],
            "Owner": "Kush",
            "Next_Action_Date": "2026-05-04",
            "Followup_1_Date": "2026-05-04",
            "Followup_2_Date": "",
            "Cold_Date": "",
            "Subject": payload["subject"],
            "Message": f"{payload['message']}\n\nRule: {payload['rule']}",
        }
        for header, value in values.items():
            if header in headers:
                ws.cell(row_idx, headers[header]).value = value
        existing[name] = row_idx


def update_execution_task(wb: Any, operating_date: date, execution_date: date, target_count: int, variant_counts: Counter) -> None:
    ws = wb["FNOMO_EXECUTION_TASKS"]
    _, headers = find_header_row(ws, "Task_ID")
    task_id = "TDA-MTB-FU1-001"
    row_idx = None
    for idx in range(1, ws.max_row + 1):
        if compact(ws.cell(idx, headers["Task_ID"]).value) == task_id:
            row_idx = idx
            break
    row_idx = row_idx or ws.max_row + 1
    values = {
        "Task_ID": task_id,
        "Task": "Messaging Test Batch Follow-up 1",
        "Owner": "Kush",
        "Priority": "A",
        "Next_Action": "Execute Monday Follow-up 1 for Tier A BO/CA Messaging Test Batch using variant rules; no links, no product, no pitch.",
        "Deadline": execution_date.isoformat(),
        "Status": "Planned for Monday",
        "Related_Segment": "Messaging Test Batch 1",
        "Notes": f"Prepared on {operating_date.isoformat()}; targets={target_count}; variants={dict(variant_counts)}; Sunday planning only.",
    }
    for header, value in values.items():
        if header in headers:
            ws.cell(row_idx, headers[header]).value = value


def append_history(wb: Any, operating_date: date, execution_date: date, target_count: int, variant_counts: Counter, held_count: int) -> None:
    ws = wb["FNOMO_EXECUTION_HISTORY"]
    for row_idx in range(ws.max_row, 1, -1):
        if (
            compact(ws.cell(row_idx, 1).value) == operating_date.isoformat()
            and compact(ws.cell(row_idx, 2).value) == "Follow-up 1 Variant Planning"
            and compact(ws.cell(row_idx, 3).value) == "Messaging Test Batch Tier A BO/CA"
        ):
            ws.delete_rows(row_idx, 1)
    row_idx = ws.max_row + 1
    values = [
        operating_date.isoformat(),
        "Follow-up 1 Variant Planning",
        "Messaging Test Batch Tier A BO/CA",
        "Codex",
        "Workbook + Notion",
        f"Locked Follow-up 1 variants for {target_count} Tier A BO/CA contacted/no-response leads. Variant counts: {dict(variant_counts)}. Platform holds excluded: {held_count}.",
        f"Execute on {execution_date.isoformat()} only; no Sunday sending.",
        "Main chat PA control",
    ]
    for col_idx, value in enumerate(values, 1):
        ws.cell(row_idx, col_idx).value = value


def ensure_tag_options() -> None:
    for database_id in (CLIENTS_DB, CONTACTS_DB, EXECUTION_DB):
        schema = notion("GET", f"databases/{database_id}").get("properties") or {}
        if "Tags" not in schema:
            notion("PATCH", f"databases/{database_id}", {"properties": {"Tags": {"multi_select": {}}}})
        SCHEMA_CACHE[database_id] = set((notion("GET", f"databases/{database_id}").get("properties") or {}).keys())


def apply_notion(operating_date: date, execution_date: date, touched: list[dict[str, Any]], held: list[dict[str, Any]], variant_counts: dict[str, int]) -> dict[str, Any]:
    ensure_tag_options()
    clients = query_database(CLIENTS_DB)
    contacts = query_database(CONTACTS_DB)
    client_by_company = {norm(plain(prop(client, "Company")) or title_of(client)): client for client in clients}
    counts: Counter = Counter()
    for row in touched:
        client = client_by_company.get(norm(row["company"]))
        contact = find_contact(contacts, row["name"], client)
        if client:
            update_client(client, row, operating_date)
            counts["clients_updated"] += 1
        if contact:
            update_contact(contact, row, operating_date)
            counts["contacts_updated"] += 1
        if client and contact:
            upsert_task(client, contact, row, execution_date)
            counts["tasks_upserted"] += 1
    for row in held:
        client = client_by_company.get(norm(row["company"]))
        contact = find_contact(contacts, row["name"], client)
        if client:
            update_platform_hold_client(client, row, operating_date)
            counts["platform_clients_held"] += 1
        if contact:
            update_platform_hold_contact(contact, row, operating_date)
            counts["platform_contacts_held"] += 1
        if client and contact:
            upsert_platform_hold_task(client, contact, row, execution_date)
            counts["platform_tasks_held"] += 1
    log_notion_history(operating_date, execution_date, counts, len(touched), variant_counts, len(held))
    return dict(counts)


def find_contact(contacts: list[dict[str, Any]], name: str, client: dict[str, Any] | None) -> dict[str, Any] | None:
    name_key = norm(name)
    client_id = client["id"] if client else ""
    for contact in contacts:
        if norm(plain(prop(contact, "Name")) or title_of(contact)) != name_key:
            continue
        if client_id and client_id in relation_ids(contact, "Company"):
            return contact
    for contact in contacts:
        if norm(plain(prop(contact, "Name")) or title_of(contact)) == name_key:
            return contact
    return None


def update_client(page: dict[str, Any], row: dict[str, Any], operating_date: date) -> None:
    note = f"[{operating_date.isoformat()} Sunday planning] Follow-up 1 variant: {row['variant']}."
    props = {
        "Stage": select_value("Contacted"),
        "Status": select_value("Active"),
        "Tier": select_value("A"),
        "Last Contact Date": date_value(row["last_contact"]),
        "Next Action": rich_text(row["next_action"]),
        "Next Action Date": date_value(row["next_date"]),
        "Notes": rich_text(append_note(plain(prop(page, "Notes")), note)),
        "Tags": multi_select(prop(page, "Tags"), [OPERATING_TAG, FOLLOWUP_TAG, PLANNING_TAG]),
    }
    notion("PATCH", f"pages/{page['id']}", {"properties": allowed_props(CLIENTS_DB, props)})


def update_contact(page: dict[str, Any], row: dict[str, Any], operating_date: date) -> None:
    note = f"[{operating_date.isoformat()} Sunday planning] Follow-up 1 variant: {row['variant']}."
    props = {
        "Stage": select_value("Contacted"),
        "Tier": select_value("A"),
        "Last Contact Date": date_value(row["last_contact"]),
        "Next Action": rich_text(row["next_action"]),
        "Next Action Date": date_value(row["next_date"]),
        "Notes": rich_text(append_note(plain(prop(page, "Notes")), note)),
        "Tags": multi_select(prop(page, "Tags"), [OPERATING_TAG, FOLLOWUP_TAG, PLANNING_TAG]),
    }
    notion("PATCH", f"pages/{page['id']}", {"properties": allowed_props(CONTACTS_DB, props)})


def update_platform_hold_client(page: dict[str, Any], row: dict[str, Any], operating_date: date) -> None:
    note = f"[{operating_date.isoformat()} Sunday planning] Excluded from BO/CA Follow-up 1 variants; platform route needs separate angle."
    props = {
        "Next Action": rich_text(row["next_action"]),
        "Next Action Date": date_value(row["next_date"]),
        "Notes": rich_text(append_note(strip_stale_platform_variant_note(plain(prop(page, "Notes"))), note)),
        "Tags": multi_select(prop(page, "Tags"), [OPERATING_TAG, PLATFORM_HOLD_TAG, PLANNING_TAG]),
    }
    notion("PATCH", f"pages/{page['id']}", {"properties": allowed_props(CLIENTS_DB, props)})


def update_platform_hold_contact(page: dict[str, Any], row: dict[str, Any], operating_date: date) -> None:
    note = f"[{operating_date.isoformat()} Sunday planning] Excluded from BO/CA Follow-up 1 variants; platform route needs separate angle."
    props = {
        "Next Action": rich_text(row["next_action"]),
        "Next Action Date": date_value(row["next_date"]),
        "Notes": rich_text(append_note(strip_stale_platform_variant_note(plain(prop(page, "Notes"))), note)),
        "Tags": multi_select(prop(page, "Tags"), [OPERATING_TAG, PLATFORM_HOLD_TAG, PLANNING_TAG]),
    }
    notion("PATCH", f"pages/{page['id']}", {"properties": allowed_props(CONTACTS_DB, props)})


def upsert_task(client: dict[str, Any], contact: dict[str, Any], row: dict[str, Any], execution_date: date) -> None:
    title = f"Follow-up 1 with {row['name']} ({row['company']})"
    existing = query_database(
        EXECUTION_DB,
        {
            "and": [
                {"property": "Task", "title": {"equals": title}},
                {"property": "Contact", "relation": {"contains": contact["id"]}},
            ]
        },
    )
    props = {
        "Task": title_text(title),
        "Client": {"relation": [{"id": client["id"]}]},
        "Contact": {"relation": [{"id": contact["id"]}]},
        "Type": select_value("Follow-Up"),
        "Priority": select_value("High"),
        "Due date": date_value(execution_date.isoformat()),
        "Status": status_value("To Do"),
        "Stage Context": select_value("Contacted"),
        "Additional Information": rich_text(f"{row['variant']}: {VARIANTS[row['variant']]['message']}"),
        "Tags": {"multi_select": [{"name": OPERATING_TAG}, {"name": FOLLOWUP_TAG}, {"name": PLANNING_TAG}]},
    }
    props = allowed_props(EXECUTION_DB, props)
    if existing:
        notion("PATCH", f"pages/{existing[0]['id']}", {"properties": props})
    else:
        notion("POST", "pages", {"parent": {"database_id": EXECUTION_DB}, "properties": props})


def upsert_platform_hold_task(client: dict[str, Any], contact: dict[str, Any], row: dict[str, Any], execution_date: date) -> None:
    old_title = f"Follow-up 1 with {row['name']} ({row['company']})"
    old_tasks = query_database(
        EXECUTION_DB,
        {
            "and": [
                {"property": "Task", "title": {"equals": old_title}},
                {"property": "Contact", "relation": {"contains": contact["id"]}},
            ]
        },
    )
    for old_task in old_tasks:
        notion("PATCH", f"pages/{old_task['id']}", {"archived": True})

    title = f"Prepare platform follow-up with {row['name']} ({row['company']})"
    existing = query_database(
        EXECUTION_DB,
        {
            "and": [
                {"property": "Task", "title": {"equals": title}},
                {"property": "Contact", "relation": {"contains": contact["id"]}},
            ]
        },
    )
    props = {
        "Task": title_text(title),
        "Client": {"relation": [{"id": client["id"]}]},
        "Contact": {"relation": [{"id": contact["id"]}]},
        "Type": select_value("Follow-Up"),
        "Priority": select_value("High"),
        "Due date": date_value(execution_date.isoformat()),
        "Status": status_value("To Do"),
        "Stage Context": select_value("Contacted"),
        "Additional Information": rich_text(row["next_action"]),
        "Tags": {"multi_select": [{"name": OPERATING_TAG}, {"name": PLATFORM_HOLD_TAG}, {"name": PLANNING_TAG}]},
    }
    props = allowed_props(EXECUTION_DB, props)
    if existing:
        notion("PATCH", f"pages/{existing[0]['id']}", {"properties": props})
    else:
        notion("POST", "pages", {"parent": {"database_id": EXECUTION_DB}, "properties": props})


def log_notion_history(operating_date: date, execution_date: date, counts: Counter, target_count: int, variant_counts: dict[str, int], held_count: int) -> None:
    key = f"{operating_date.isoformat()}::followup1-variants-messaging-test-batch"
    existing = query_database(HISTORY_DB, {"property": "History Key", "title": {"equals": key}})
    props = {
        "History Key": title_text(key),
        "Date": date_value(operating_date.isoformat()),
        "Action Type": rich_text("Follow-up 1 Variant Planning"),
        "Lead or Task": rich_text("Messaging Test Batch Tier A BO/CA"),
        "Owner": rich_text("Codex"),
        "Channel": rich_text("Workbook + Notion"),
        "Outcome": {"select": {"name": "No Response"}},
        "Outcome Notes": rich_text(f"Prepared Monday Follow-up 1 variants for {target_count} Tier A BO/CA contacted/no-response leads. Variant counts: {variant_counts}; platform holds excluded: {held_count}; Notion counts: {dict(counts)}."),
        "Next Move": rich_text(f"Execute on {execution_date.isoformat()}; no links, no product, no pitch. Sunday is planning only."),
        "Source": rich_text("Main chat PA control"),
        "Last Synced At": {"date": {"start": datetime.now().astimezone().isoformat(timespec="seconds")}},
    }
    if existing:
        notion("PATCH", f"pages/{existing[0]['id']}", {"properties": props})
    else:
        notion("POST", "pages", {"parent": {"database_id": HISTORY_DB}, "properties": props})


def validate_workbook() -> dict[str, Any]:
    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    pipeline = wb["FNOMO_MASTER_PIPELINE"]
    _, headers = find_header_row(pipeline, "Lead_ID")
    rows = read_pipeline_rows(pipeline, headers)
    targets = [row for row in rows if target_row(row)]
    held = [row for row in rows if platform_hold_row(row) and not target_row(row)]
    return {
        "target_rows": len(targets),
        "target_due_dates": dict(Counter(row["next_date"] for row in targets)),
        "target_actions_missing_followup": sum(1 for row in targets if "Follow-up 1" not in row["next_action"]),
        "platform_holds": len(held),
    }


def main() -> None:
    operating_date = date.fromisoformat(os.environ.get("FNOMO_OPERATING_DATE", date.today().isoformat()))
    execution_date = date.fromisoformat(os.environ.get("FNOMO_EXECUTION_DATE", "2026-05-04"))
    workbook_result = apply_workbook(operating_date, execution_date)
    notion_result = apply_notion(operating_date, execution_date, workbook_result["touched"], workbook_result["held"], workbook_result["variant_counts"])
    validation = validate_workbook()
    print(json.dumps({"workbook": workbook_result, "notion": notion_result, "validation": validation}, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
