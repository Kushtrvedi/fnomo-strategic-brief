#!/usr/bin/env python3
"""
Repair and sync the latest Tier A priority follow-up update into Notion.

Source of truth: FNOMO_MASTER_PIPELINE in the master workbook.
Scope is intentionally narrow: May 4 priority BO/CA follow-up correction.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import time
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from shutil import copy2
from typing import Any

import requests
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "EXCEL" / "War Room_ Community Outreach Pipeline - FNOMO Master.xlsx"
CLIENTS_DB = "e8a8aa3c-b995-8368-8495-014b8ea295c3"
CONTACTS_DB = "3548aa3c-b995-8148-a653-fead444e56fe"
EXECUTION_DB = "6fb8aa3c-b995-82ec-806b-01b6cd1b2915"
HISTORY_DB = "3548aa3c-b995-818a-8ede-ed9a5d99aac3"
NOTION_VERSION = "2022-06-28"
SESSION = requests.Session()
SCHEMA_CACHE: dict[str, set[str]] = {}

SYNC_TAG = "Priority Follow-up Sync"
FOLLOWUP_TAG = "Follow-up 1"


TARGETS: dict[tuple[str, str], dict[str, Any]] = {
    ("Nitin Nahar", "Nahar Enterprise"): {
        "stage": "Contacted",
        "status": "Active",
        "tier": "A",
        "next_date": "2026-05-04",
        "last_contact": "2026-05-02",
        "next_action": "Follow-up 1: reopen with the decision-validation question; ask what capital decision looked right but should have been stress-tested before money moved.",
    },
    ("Sandeep Wadhwani", "Sandeep Wadhwani"): {
        "stage": "Contacted",
        "status": "Active",
        "tier": "A",
        "next_date": "2026-05-04",
        "last_contact": "2026-05-02",
        "next_action": "Follow-up 1: reopen with the decision-validation question; ask what business decision looked right but should have been stress-tested before money moved.",
    },
    ("Sumit Lalwani", "Sumit Lalwani"): {
        "stage": "Contacted",
        "status": "Active",
        "tier": "A",
        "next_date": "2026-05-04",
        "last_contact": "2026-05-02",
        "next_action": "Follow-up 1: reopen with the decision-validation question around stock, supplier, cash-cycle, or expansion risk before money moves.",
    },
    ("Kishore Gupta", "Star Delta Transformers Ltd"): {
        "stage": "Contacted",
        "status": "Active",
        "tier": "A",
        "next_date": "2026-05-04",
        "last_contact": "2026-05-02",
        "next_action": "Follow-up 1: reopen with the decision-validation question around capex, procurement, customer-credit, or production risk before money moves.",
    },
    ("Vinod Sapre", "Righill Electrics Pvt. Ltd."): {
        "stage": "Contacted",
        "status": "Active",
        "tier": "A",
        "next_date": "2026-05-04",
        "last_contact": "2026-05-02",
        "next_action": "Follow-up 1: reopen with the decision-validation question around industrial control, vendor, customer, or capacity assumptions before money moves.",
    },
    ("K.M. Kumar", "Kumaran Industries"): {
        "stage": "Contacted",
        "status": "Active",
        "tier": "A",
        "next_date": "2026-05-04",
        "last_contact": "2026-05-02",
        "next_action": "Follow-up 1: reopen with the decision-validation question around quote, spec, vendor, delivery, or capital commitment risk before money moves.",
    },
    ("Milind Wadhwani", "Indore Branch of CIRC of ICAI"): {
        "stage": "Contacted",
        "status": "Active",
        "tier": "A",
        "next_date": "2026-05-04",
        "last_contact": "2026-05-02",
        "next_action": "Follow-up 1: ask whether any client made an investment decision that went wrong, then route an Indore member conversation through Milind/branch leadership.",
    },
    ("Rajat Dhanuka", "Indore Branch of CIRC of ICAI"): {
        "stage": "New",
        "status": "Waiting",
        "tier": "B",
        "next_date": "2026-05-05",
        "last_contact": None,
        "next_action": "Hold Rajat as prior-chair/support route. Do not follow up as contacted unless a direct send is logged; use only if Samkit/Mausam/Milind ownership routing stalls.",
    },
}


AGGREGATE_TARGETS: dict[str, dict[str, Any]] = {
    "Business Owners >2Cr": {
        "next_date": "2026-05-04",
        "next_action": "2026-05-04 mandate: complete Follow-up 1 for contacted Tier A owners first (Nahar, Deep, Tirupati, Star Delta, Righill, Kumaran); then send the verified unsent owner/CMD batch only in the listed order.",
    },
    "CA Associations": {
        "next_date": "2026-05-04",
        "next_action": "2026-05-04 mandate: log Arpit/Abhishek Follow-up 1, then Milind follow-up/Indore routing; use Samank/Mausam/Samkit as ownership checks before any large-branch expansion.",
    },
}


def die(message: str) -> None:
    print(f"ERROR: {message}", file=sys.stderr)
    raise SystemExit(1)


def headers() -> dict[str, str]:
    token = os.environ.get("NOTION_API_KEY")
    if not token:
        die("NOTION_API_KEY missing")
    return {
        "Authorization": f"Bearer {token}",
        "Notion-Version": NOTION_VERSION,
        "Content-Type": "application/json",
    }


def notion(method: str, path: str, payload: dict[str, Any] | None = None) -> dict[str, Any]:
    for attempt in range(6):
        response = SESSION.request(
            method,
            f"https://api.notion.com/v1/{path.lstrip('/')}",
            headers=headers(),
            json=payload,
            timeout=60,
        )
        if response.status_code == 429:
            time.sleep(2 + attempt)
            continue
        if response.status_code >= 400:
            raise RuntimeError(f"{method} {path} failed {response.status_code}: {response.text[:1500]}")
        return response.json()
    raise RuntimeError(f"{method} {path} failed after retries")


def query_database(database_id: str, filter_payload: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    cursor = None
    while True:
        payload: dict[str, Any] = {"page_size": 100}
        if cursor:
            payload["start_cursor"] = cursor
        if filter_payload:
            payload["filter"] = filter_payload
        data = notion("POST", f"databases/{database_id}/query", payload)
        rows.extend(data.get("results", []))
        if not data.get("has_more"):
            return rows
        cursor = data.get("next_cursor")


def compact(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    return str(value).strip()


def norm(value: Any) -> str:
    return " ".join(compact(value).casefold().split())


def plain(prop: dict[str, Any] | None) -> str:
    if not prop:
        return ""
    prop_type = prop.get("type")
    if prop_type in {"title", "rich_text"}:
        return "".join(part.get("plain_text", "") for part in prop.get(prop_type, [])).strip()
    if prop_type == "select":
        return ((prop.get("select") or {}).get("name") or "").strip()
    if prop_type == "status":
        return ((prop.get("status") or {}).get("name") or "").strip()
    if prop_type == "date" and prop.get("date"):
        return (prop["date"].get("start") or "").strip()
    if prop_type == "multi_select":
        return "; ".join(item.get("name", "") for item in prop.get("multi_select", []))
    return ""


def prop(page: dict[str, Any], name: str) -> dict[str, Any] | None:
    return (page.get("properties") or {}).get(name)


def title_of(page: dict[str, Any]) -> str:
    for value in (page.get("properties") or {}).values():
        if value.get("type") == "title":
            return plain(value)
    return ""


def page_name(page: dict[str, Any]) -> str:
    return plain(prop(page, "Name")) or plain(prop(page, "Company")) or title_of(page)


def title_text(value: str) -> dict[str, Any]:
    return {"title": [{"text": {"content": value[:1800]}}]}


def rich_text(value: Any) -> dict[str, Any]:
    text = compact(value)[:1900]
    return {"rich_text": [{"text": {"content": text}}]} if text else {"rich_text": []}


def date_value(value: Any) -> dict[str, Any]:
    text = compact(value)
    return {"date": {"start": text}} if text else {"date": None}


def select_value(value: Any) -> dict[str, Any]:
    text = compact(value)[:100]
    return {"select": {"name": text}} if text else {"select": None}


def status_value(value: Any) -> dict[str, Any]:
    text = compact(value)[:100]
    return {"status": {"name": text}} if text else {"status": None}


def multi_select(existing_prop: dict[str, Any] | None, tags: list[str]) -> dict[str, Any]:
    existing = []
    if existing_prop and existing_prop.get("type") == "multi_select":
        existing = [item.get("name", "") for item in existing_prop.get("multi_select", []) if item.get("name")]
    for tag in tags:
        if tag and tag not in existing:
            existing.append(tag)
    return {"multi_select": [{"name": tag} for tag in existing]}


def append_note(existing: Any, note: str) -> str:
    text = compact(existing)
    if note in text:
        return text
    return (text + " | " + note).strip(" |")


def find_header_row(ws: Any, required: str) -> tuple[int, dict[str, int]]:
    for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
        headers = [compact(cell) for cell in row]
        if required in headers:
            return row_index, {header: idx + 1 for idx, header in enumerate(headers) if header}
    die(f"Missing header {required}")


def apply_workbook(today: date) -> dict[str, Any]:
    backup = WORKBOOK.parent / "backups" / f"{WORKBOOK.stem}.before-mainchat-priority-notion-sync.{datetime.now().strftime('%Y%m%d-%H%M%S')}{WORKBOOK.suffix}"
    backup.parent.mkdir(exist_ok=True)
    copy2(WORKBOOK, backup)
    wb = load_workbook(WORKBOOK)
    ws = wb["FNOMO_MASTER_PIPELINE"]
    _, headers = find_header_row(ws, "Lead_ID")
    touched: list[dict[str, Any]] = []
    summary_note = f"[{today.isoformat()} main-chat sync] May 4 priority follow-up synced to Notion; sheet remains source of truth."

    for row_idx in range(1, ws.max_row + 1):
        name = compact(ws.cell(row_idx, headers["Name"]).value)
        company = compact(ws.cell(row_idx, headers["Company"]).value)
        key = (name, company)
        target = TARGETS.get(key)
        if target:
            ws.cell(row_idx, headers["Tier (A/B/C)"]).value = target["tier"]
            ws.cell(row_idx, headers["Stage (New / Contacted / Engaged / Qualified / Converted / Lost)"]).value = target["stage"]
            ws.cell(row_idx, headers["Status (Active / Waiting / No Response / Closed)"]).value = target["status"]
            ws.cell(row_idx, headers["Last_Contact_Date"]).value = target["last_contact"]
            ws.cell(row_idx, headers["Next_Action"]).value = target["next_action"]
            ws.cell(row_idx, headers["Next_Action_Date"]).value = target["next_date"]
            ws.cell(row_idx, headers["Notes (free text)"]).value = append_note(ws.cell(row_idx, headers["Notes (free text)"]).value, summary_note)
            tags = compact(ws.cell(row_idx, headers.get("Tags", 0)).value) if "Tags" in headers else ""
            if "Tags" in headers:
                ws.cell(row_idx, headers["Tags"]).value = "; ".join(dict.fromkeys([t.strip() for t in (tags + "; " + SYNC_TAG).split(";") if t.strip()]))
            touched.append({"name": name, "company": company, **target})
        elif not name and company in AGGREGATE_TARGETS:
            target = AGGREGATE_TARGETS[company]
            ws.cell(row_idx, headers["Next_Action"]).value = target["next_action"]
            ws.cell(row_idx, headers["Next_Action_Date"]).value = target["next_date"]
            ws.cell(row_idx, headers["Notes (free text)"]).value = append_note(ws.cell(row_idx, headers["Notes (free text)"]).value, summary_note)

    append_history(wb, today)
    wb.save(WORKBOOK)
    return {"backup": str(backup), "touched": touched}


def append_history(wb: Any, today: date) -> None:
    ws = wb["FNOMO_EXECUTION_HISTORY"]
    for row_index in range(ws.max_row, 1, -1):
        if (
            compact(ws.cell(row_index, 1).value) == today.isoformat()
            and compact(ws.cell(row_index, 2).value) == "Main Chat Notion Sync"
            and compact(ws.cell(row_index, 3).value) == "May 4 Tier A priority follow-ups"
        ):
            ws.delete_rows(row_index, 1)
    row = ws.max_row + 1
    values = [
        today.isoformat(),
        "Main Chat Notion Sync",
        "May 4 Tier A priority follow-ups",
        "Codex",
        "Workbook + Notion",
        "Repaired individual Tier A BO/CA follow-up rows to May 4 and synced the priority update into Notion.",
        "Execute May 4 Follow-up 1 first; then log sends/replies back into workbook before broader expansion.",
        "Main chat PA control",
    ]
    for col_idx, value in enumerate(values, start=1):
        ws.cell(row, col_idx).value = value


def ensure_tag_options() -> None:
    for db in (CLIENTS_DB, CONTACTS_DB, EXECUTION_DB):
        schema = notion("GET", f"databases/{db}").get("properties") or {}
        if "Tags" not in schema:
            notion("PATCH", f"databases/{db}", {"properties": {"Tags": {"multi_select": {}}}})
        SCHEMA_CACHE[db] = set((notion("GET", f"databases/{db}").get("properties") or {}).keys())


def allowed_props(database_id: str, props: dict[str, Any]) -> dict[str, Any]:
    if database_id not in SCHEMA_CACHE:
        SCHEMA_CACHE[database_id] = set((notion("GET", f"databases/{database_id}").get("properties") or {}).keys())
    return {key: value for key, value in props.items() if key in SCHEMA_CACHE[database_id]}


def apply_notion(today: date, rows: list[dict[str, Any]]) -> dict[str, Any]:
    ensure_tag_options()
    contacts = query_database(CONTACTS_DB)
    clients = query_database(CLIENTS_DB)
    contact_by_key = {(norm(plain(prop(c, "Name")) or title_of(c)), norm(plain(prop(c, "Company")) or "")): c for c in contacts}
    client_by_company = {norm(plain(prop(c, "Company")) or title_of(c)): c for c in clients}
    counts: Counter = Counter()
    for row in rows:
        contact = find_contact(contact_by_key, row["name"], row["company"])
        client = client_by_company.get(norm(row["company"]))
        if contact:
            update_contact(contact, row, today)
            counts["contacts_updated"] += 1
        if client:
            update_client(client, row, today)
            counts["clients_updated"] += 1
        if client and contact and row["stage"] == "Contacted":
            upsert_task(client, contact, row, today)
            counts["tasks_upserted"] += 1
    log_notion_history(today, counts, rows)
    return dict(counts)


def find_contact(contact_by_key: dict[tuple[str, str], dict[str, Any]], name: str, company: str) -> dict[str, Any] | None:
    key = (norm(name), norm(company))
    if key in contact_by_key:
        return contact_by_key[key]
    for (candidate_name, candidate_company), page in contact_by_key.items():
        if candidate_name == norm(name) and (candidate_company == norm(company) or not candidate_company):
            return page
    return None


def update_contact(page: dict[str, Any], row: dict[str, Any], today: date) -> None:
    note = f"[{today.isoformat()} main-chat sync] {row['next_action']}"
    props = {
        "Stage": select_value(row["stage"]),
        "Status": select_value(row["status"]),
        "Tier": select_value(row["tier"]),
        "Last Contact Date": date_value(row["last_contact"]),
        "Next Action": rich_text(row["next_action"]),
        "Next Action Date": date_value(row["next_date"]),
        "Notes": rich_text(append_note(plain(prop(page, "Notes")), note)),
        "Tags": multi_select(prop(page, "Tags"), [SYNC_TAG, FOLLOWUP_TAG] if row["stage"] == "Contacted" else [SYNC_TAG, "Support Route"]),
    }
    notion("PATCH", f"pages/{page['id']}", {"properties": allowed_props(CONTACTS_DB, props)})


def update_client(page: dict[str, Any], row: dict[str, Any], today: date) -> None:
    note = f"[{today.isoformat()} main-chat sync] {row['name']}: {row['next_action']}"
    props = {
        "Stage": select_value(row["stage"]),
        "Status": select_value(row["status"] if row["stage"] != "New" else "Lead"),
        "Tier": select_value(row["tier"]),
        "Last Contact Date": date_value(row["last_contact"]),
        "Next Action": rich_text(row["next_action"]),
        "Next Action Date": date_value(row["next_date"]),
        "Notes": rich_text(append_note(plain(prop(page, "Notes")), note)),
        "Tags": multi_select(prop(page, "Tags"), [SYNC_TAG, FOLLOWUP_TAG] if row["stage"] == "Contacted" else [SYNC_TAG, "Support Route"]),
    }
    notion("PATCH", f"pages/{page['id']}", {"properties": allowed_props(CLIENTS_DB, props)})


def upsert_task(client: dict[str, Any], contact: dict[str, Any], row: dict[str, Any], today: date) -> None:
    title = f"Follow-up 1 with {row['name']} ({row['company']})"
    existing = query_database(
        EXECUTION_DB,
        {
            "and": [
                {"property": "Task", "title": {"equals": title}},
                {"property": "Contact", "relation": {"contains": contact["id"]}},
            ]
        },
    )
    props = {
        "Task": title_text(title),
        "Client": {"relation": [{"id": client["id"]}]},
        "Contact": {"relation": [{"id": contact["id"]}]},
        "Type": select_value("Follow-Up"),
        "Priority": select_value("High"),
        "Due date": date_value(row["next_date"]),
        "Status": status_value("To Do"),
        "Stage Context": select_value(row["stage"]),
        "Additional Information": rich_text(row["next_action"]),
        "Tags": {"multi_select": [{"name": SYNC_TAG}, {"name": FOLLOWUP_TAG}]},
    }
    props = allowed_props(EXECUTION_DB, props)
    if existing:
        notion("PATCH", f"pages/{existing[0]['id']}", {"properties": props})
    else:
        notion("POST", "pages", {"parent": {"database_id": EXECUTION_DB}, "properties": props})


def log_notion_history(today: date, counts: Counter, rows: list[dict[str, Any]]) -> None:
    key = f"{today.isoformat()}::main-chat-may4-priority-followups-notion-sync"
    existing = query_database(HISTORY_DB, {"property": "History Key", "title": {"equals": key}})
    properties = {
        "History Key": title_text(key),
        "Date": date_value(today.isoformat()),
        "Action Type": rich_text("Main Chat Notion Sync"),
        "Lead or Task": rich_text("May 4 Tier A priority follow-ups"),
        "Owner": rich_text("Codex"),
        "Channel": rich_text("Workbook + Notion"),
        "Outcome": {"select": {"name": "No Response"}},
        "Outcome Notes": rich_text(f"Synced May 4 follow-up priority correction into Notion. Counts: {dict(counts)}. Rows: {', '.join(r['name'] for r in rows)}."),
        "Next Move": rich_text("Execute May 4 Follow-up 1 first; then log sends/replies before broader expansion."),
        "Source": rich_text("Main chat PA control"),
        "Last Synced At": {"date": {"start": datetime.now().astimezone().isoformat(timespec="seconds")}},
    }
    if existing:
        notion("PATCH", f"pages/{existing[0]['id']}", {"properties": properties})
    else:
        notion("POST", "pages", {"parent": {"database_id": HISTORY_DB}, "properties": properties})


def validate_workbook() -> dict[str, Any]:
    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    ws = wb["FNOMO_MASTER_PIPELINE"]
    _, headers = find_header_row(ws, "Lead_ID")
    result = {}
    for row in ws.iter_rows(values_only=True):
        name = compact(row[headers["Name"] - 1] if headers["Name"] <= len(row) else "")
        company = compact(row[headers["Company"] - 1] if headers["Company"] <= len(row) else "")
        if (name, company) in TARGETS:
            result[f"{name}::{company}"] = {
                "stage": compact(row[headers["Stage (New / Contacted / Engaged / Qualified / Converted / Lost)"] - 1]),
                "next_action_date": compact(row[headers["Next_Action_Date"] - 1]),
                "next_action": compact(row[headers["Next_Action"] - 1])[:120],
            }
    return result


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Repair and optionally sync the May 4 Tier A priority follow-up update into Notion."
    )
    parser.add_argument(
        "--date",
        default=os.environ.get("FNOMO_OPERATING_DATE", date.today().isoformat()),
        help="Operating date for workbook/history notes. Defaults to FNOMO_OPERATING_DATE or today.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Validate current workbook target rows without writing to the workbook or Notion.",
    )
    parser.add_argument(
        "--workbook-only",
        action="store_true",
        help="Apply workbook correction and skip Notion writes.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    today = date.fromisoformat(args.date)
    if args.dry_run:
        print(json.dumps({"dry_run": True, "validation": validate_workbook()}, indent=2, ensure_ascii=False))
        return

    workbook_result = apply_workbook(today)
    if args.workbook_only:
        validation = validate_workbook()
        print(json.dumps({"workbook": workbook_result, "notion": "skipped", "validation": validation}, indent=2, ensure_ascii=False))
        return

    notion_result = apply_notion(today, workbook_result["touched"])
    validation = validate_workbook()
    print(json.dumps({"workbook": workbook_result, "notion": notion_result, "validation": validation}, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
