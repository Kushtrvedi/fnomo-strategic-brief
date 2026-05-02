#!/usr/bin/env python3
"""
Apply the Fnomo Sunday planning-only rule.

If the operating date is Sunday, roll live next-action/task dates from Sunday
to Monday in the workbook and Notion so execution does not show false overdue.
Sunday can be used for planning and CRM preparation, not outbound execution.
"""

from __future__ import annotations

import json
import os
import sys
import time
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path
from shutil import copy2
from typing import Any

import requests
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "EXCEL" / "War Room_ Community Outreach Pipeline - FNOMO Master.xlsx"
CLIENTS_DB = "e8a8aa3c-b995-8368-8495-014b8ea295c3"
CONTACTS_DB = "3548aa3c-b995-8148-a653-fead444e56fe"
EXECUTION_DB = "6fb8aa3c-b995-82ec-806b-01b6cd1b2915"
HISTORY_DB = "3548aa3c-b995-818a-8ede-ed9a5d99aac3"
NOTION_VERSION = "2022-06-28"
SESSION = requests.Session()


def die(message: str) -> None:
    print(f"ERROR: {message}", file=sys.stderr)
    raise SystemExit(1)


def compact(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    return str(value).strip()


def parse_date(value: Any) -> str:
    text = compact(value)
    if not text:
        return ""
    if len(text) >= 10 and text[4] == "-" and text[7] == "-":
        return text[:10]
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            pass
    return text


def headers() -> dict[str, str]:
    token = os.environ.get("NOTION_API_KEY")
    if not token:
        die("NOTION_API_KEY missing")
    return {
        "Authorization": f"Bearer {token}",
        "Notion-Version": NOTION_VERSION,
        "Content-Type": "application/json",
    }


def notion(method: str, path: str, payload: dict[str, Any] | None = None) -> dict[str, Any]:
    for attempt in range(6):
        response = SESSION.request(
            method,
            f"https://api.notion.com/v1/{path.lstrip('/')}",
            headers=headers(),
            json=payload,
            timeout=60,
        )
        if response.status_code == 429:
            time.sleep(2 + attempt)
            continue
        if response.status_code >= 400:
            raise RuntimeError(f"{method} {path} failed {response.status_code}: {response.text[:1500]}")
        return response.json()
    raise RuntimeError(f"{method} {path} failed after retries")


def query_database(database_id: str, filter_payload: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    cursor = None
    while True:
        payload: dict[str, Any] = {"page_size": 100}
        if cursor:
            payload["start_cursor"] = cursor
        if filter_payload:
            payload["filter"] = filter_payload
        data = notion("POST", f"databases/{database_id}/query", payload)
        rows.extend(data.get("results", []))
        if not data.get("has_more"):
            return rows
        cursor = data.get("next_cursor")


def plain(prop: dict[str, Any] | None) -> str:
    if not prop:
        return ""
    prop_type = prop.get("type")
    if prop_type in {"title", "rich_text"}:
        return "".join(part.get("plain_text", "") for part in prop.get(prop_type, [])).strip()
    if prop_type == "date" and prop.get("date"):
        return (prop["date"].get("start") or "").strip()
    if prop_type == "select":
        return ((prop.get("select") or {}).get("name") or "").strip()
    if prop_type == "status":
        return ((prop.get("status") or {}).get("name") or "").strip()
    return ""


def prop(page: dict[str, Any], name: str) -> dict[str, Any] | None:
    return (page.get("properties") or {}).get(name)


def rich_text(value: str) -> dict[str, Any]:
    return {"rich_text": [{"text": {"content": value[:1900]}}]} if value else {"rich_text": []}


def date_value(value: str) -> dict[str, Any]:
    return {"date": {"start": value}} if value else {"date": None}


def append_note(existing: Any, note: str) -> str:
    text = compact(existing)
    if note in text:
        return text
    return (text + " | " + note).strip(" |")


def find_header_row(ws: Any, required: str) -> tuple[int, dict[str, int]]:
    for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
        headers = [compact(cell) for cell in row]
        if required in headers:
            return row_index, {header: idx + 1 for idx, header in enumerate(headers) if header}
    die(f"Missing header {required}")


def apply_workbook(operating_date: date, next_workday: date) -> dict[str, Any]:
    backup = WORKBOOK.parent / "backups" / f"{WORKBOOK.stem}.before-no-sunday-roll.{datetime.now().strftime('%Y%m%d-%H%M%S')}{WORKBOOK.suffix}"
    backup.parent.mkdir(exist_ok=True)
    copy2(WORKBOOK, backup)
    wb = load_workbook(WORKBOOK)
    counts: Counter = Counter()
    note = f"[{operating_date.isoformat()} Sunday planning rule] Sunday is planning-only; moved execution due action to {next_workday.isoformat()}."

    pipeline = wb["FNOMO_MASTER_PIPELINE"]
    _, headers = find_header_row(pipeline, "Lead_ID")
    for row_idx in range(1, pipeline.max_row + 1):
        cell = pipeline.cell(row_idx, headers["Next_Action_Date"])
        if parse_date(cell.value) == operating_date.isoformat():
            cell.value = next_workday.isoformat()
            notes_cell = pipeline.cell(row_idx, headers["Notes (free text)"])
            notes_cell.value = append_note(notes_cell.value, note)
            counts["pipeline_rows_moved"] += 1

    tasks = wb["FNOMO_EXECUTION_TASKS"]
    _, task_headers = find_header_row(tasks, "Task_ID")
    if "Due_Date" in task_headers:
        due_header = "Due_Date"
    elif "Due Date" in task_headers:
        due_header = "Due Date"
    else:
        due_header = "Deadline"
    for row_idx in range(1, tasks.max_row + 1):
        if due_header in task_headers and parse_date(tasks.cell(row_idx, task_headers[due_header]).value) == operating_date.isoformat():
            tasks.cell(row_idx, task_headers[due_header]).value = next_workday.isoformat()
            if "Status" in task_headers and compact(tasks.cell(row_idx, task_headers["Status"]).value) == "Execute Today":
                tasks.cell(row_idx, task_headers["Status"]).value = "Planned for Monday"
            counts["task_rows_moved"] += 1

    append_history(wb, operating_date, next_workday, counts)
    wb.save(WORKBOOK)
    return {"backup": str(backup), "counts": dict(counts)}


def append_history(wb: Any, operating_date: date, next_workday: date, counts: Counter) -> None:
    ws = wb["FNOMO_EXECUTION_HISTORY"]
    for row_index in range(ws.max_row, 1, -1):
        if (
            compact(ws.cell(row_index, 1).value) == operating_date.isoformat()
            and compact(ws.cell(row_index, 2).value) == "Sunday Planning Rule"
            and compact(ws.cell(row_index, 3).value) == "Sunday execution planned for Monday"
        ):
            ws.delete_rows(row_index, 1)
    row = ws.max_row + 1
    values = [
        operating_date.isoformat(),
        "Sunday Planning Rule",
        "Sunday execution planned for Monday",
        "Codex",
        "Workbook + Notion",
        f"Moved Sunday-due execution to {next_workday.isoformat()} while keeping Sunday as planning-only. Counts: {dict(counts)}.",
        "Use Sunday for next-day planning only; resume outbound execution Monday unless Kush explicitly overrides.",
        "Main chat PA control",
    ]
    for index, value in enumerate(values, 1):
        ws.cell(row, index).value = value


def patch_date_for_db(database_id: str, date_prop: str, note_prop: str, operating_date: date, next_workday: date) -> int:
    pages = query_database(database_id, {"property": date_prop, "date": {"equals": operating_date.isoformat()}})
    note = f"[{operating_date.isoformat()} Sunday planning rule] Sunday is planning-only; moved execution due action to {next_workday.isoformat()}."
    updated = 0
    for page in pages:
        props: dict[str, Any] = {date_prop: date_value(next_workday.isoformat())}
        if note_prop in (page.get("properties") or {}):
            props[note_prop] = rich_text(append_note(plain(prop(page, note_prop)), note))
        notion("PATCH", f"pages/{page['id']}", {"properties": props})
        updated += 1
    return updated


def patch_execution_tasks(operating_date: date, next_workday: date) -> int:
    pages = query_database(EXECUTION_DB, {"property": "Due date", "date": {"equals": operating_date.isoformat()}})
    updated = 0
    for page in pages:
        props: dict[str, Any] = {"Due date": date_value(next_workday.isoformat())}
        if "Additional Information" in (page.get("properties") or {}):
            props["Additional Information"] = rich_text(
                append_note(
                    plain(prop(page, "Additional Information")),
                    f"[{operating_date.isoformat()} Sunday planning rule] Sunday is planning-only; due date moved to {next_workday.isoformat()}.",
                )
            )
        notion("PATCH", f"pages/{page['id']}", {"properties": props})
        updated += 1
    return updated


def log_notion_history(operating_date: date, next_workday: date, counts: dict[str, Any]) -> None:
    key = f"{operating_date.isoformat()}::sunday-planning-only-rule"
    existing = query_database(HISTORY_DB, {"property": "History Key", "title": {"equals": key}})
    properties = {
        "History Key": {"title": [{"text": {"content": key}}]},
        "Date": date_value(operating_date.isoformat()),
        "Action Type": rich_text("Sunday Planning Rule"),
        "Lead or Task": rich_text("Sunday execution planned for Monday"),
        "Owner": rich_text("Codex"),
        "Channel": rich_text("Workbook + Notion"),
        "Outcome": {"select": {"name": "No Response"}},
        "Outcome Notes": rich_text(f"Sunday is planning-only. Moved due execution to {next_workday.isoformat()}. Counts: {counts}."),
        "Next Move": rich_text("Use Sunday for planning and Monday queue prep only; resume outbound execution Monday unless Kush explicitly overrides."),
        "Source": rich_text("Main chat PA control"),
        "Last Synced At": {"date": {"start": datetime.now().astimezone().isoformat(timespec="seconds")}},
    }
    if existing:
        notion("PATCH", f"pages/{existing[0]['id']}", {"properties": properties})
    else:
        notion("POST", "pages", {"parent": {"database_id": HISTORY_DB}, "properties": properties})


def apply_notion(operating_date: date, next_workday: date) -> dict[str, Any]:
    counts = {
        "clients_moved": patch_date_for_db(CLIENTS_DB, "Next Action Date", "Notes", operating_date, next_workday),
        "contacts_moved": patch_date_for_db(CONTACTS_DB, "Next Action Date", "Notes", operating_date, next_workday),
        "tasks_moved": patch_execution_tasks(operating_date, next_workday),
    }
    log_notion_history(operating_date, next_workday, counts)
    return counts


def main() -> None:
    operating_date = date.fromisoformat(os.environ.get("FNOMO_OPERATING_DATE", date.today().isoformat()))
    if operating_date.weekday() != 6:
        print(json.dumps({"skipped": True, "reason": f"{operating_date.isoformat()} is not Sunday"}))
        return
    next_workday = operating_date + timedelta(days=1)
    workbook_result = apply_workbook(operating_date, next_workday)
    notion_result = apply_notion(operating_date, next_workday)
    print(json.dumps({"workbook": workbook_result, "notion": notion_result}, indent=2))


if __name__ == "__main__":
    main()
