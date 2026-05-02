#!/usr/bin/env python3
"""
Mirror FNOMO_MASTER_PIPELINE and FNOMO_EXECUTION_HISTORY into Notion.

Auth: Notion internal integration token via NOTION_API_KEY.
No OAuth flow is used.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urlencode
from urllib.request import urlopen

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "EXCEL" / "War Room_ Community Outreach Pipeline - FNOMO Master.xlsx"
NOTION_VERSION = "2022-06-28"
MAX_RICH_TEXT = 1900
CONFIG_PATH = ROOT / "data" / "fnomo_notion_backend.json"


LEAD_FIELD_ALIASES = {
    "lead_id": ["Lead_ID", "Lead ID", "Lead Id"],
    "name": ["Name"],
    "company": ["Company"],
    "lead_type": ["Type (CA / Business / Association)", "Type"],
    "tier": ["Tier (A/B/C)", "Tier"],
    "stage": ["Stage (New / Contacted / Engaged / Qualified / Converted / Lost)", "Stage"],
    "last_contact_date": ["Last_Contact_Date", "Last Contact Date"],
    "next_action": ["Next_Action", "Next Action"],
    "next_action_date": ["Next_Action_Date", "Next Action Date"],
    "pipeline_status": ["Status (Active / Waiting / No Response / Closed)", "Status"],
    "notes": ["Notes (free text)", "Notes"],
}

HISTORY_FIELD_ALIASES = {
    "date": ["Date"],
    "action_type": ["Action_Type", "Action Type"],
    "lead_or_task": ["Lead_or_Task", "Lead or Task"],
    "owner": ["Owner"],
    "channel": ["Channel"],
    "outcome": ["Outcome"],
    "next_move": ["Next_Move", "Next Move"],
    "source": ["Source"],
}


def die(message: str) -> None:
    print(f"ERROR: {message}", file=sys.stderr)
    raise SystemExit(1)


def notion_headers() -> dict[str, str]:
    token = os.environ.get("NOTION_API_KEY")
    if not token:
        die("NOTION_API_KEY is missing. Set the internal integration token first.")
    return {
        "Authorization": f"Bearer {token}",
        "Notion-Version": NOTION_VERSION,
        "Content-Type": "application/json",
    }


def notion_request(method: str, path: str, payload: dict[str, Any] | None = None) -> dict[str, Any]:
    import requests

    url = f"https://api.notion.com/v1/{path.lstrip('/')}"
    response = requests.request(method, url, headers=notion_headers(), json=payload, timeout=30)
    if response.status_code >= 400:
        detail = response.text[:1000]
        raise RuntimeError(f"Notion {method} {path} failed: {response.status_code} {detail}")
    return response.json()


def normalize_header(value: Any) -> str:
    return str(value or "").strip()


def find_header_row(rows: list[list[Any]], required: str) -> int:
    for idx, row in enumerate(rows):
        if required in [normalize_header(cell) for cell in row]:
            return idx
    die(f"Could not find header row containing {required!r}.")


def value_by_alias(row: dict[str, Any], aliases: list[str]) -> Any:
    for alias in aliases:
        if alias in row:
            return row.get(alias)
    return None


def compact(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    return str(value).strip()


def load_backend_config() -> dict[str, Any]:
    if not CONFIG_PATH.exists():
        return {}
    with CONFIG_PATH.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def sync_key(name: Any, company: Any) -> str:
    left = " ".join(compact(name).casefold().split())
    right = " ".join(compact(company).casefold().split())
    return f"{left}::{right}"


def parse_date(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            pass
    if len(text) >= 10 and text[4] == "-" and text[7] == "-":
        return text[:10]
    return None


def read_xlsx_sheet(path: Path, sheet_name: str, required_header: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        die(f"Workbook missing sheet {sheet_name!r}: {path}")
    sheet = workbook[sheet_name]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    header_index = find_header_row(rows, required_header)
    headers = [normalize_header(cell) for cell in rows[header_index]]
    records: list[dict[str, Any]] = []
    for raw in rows[header_index + 1 :]:
        if not any(cell not in (None, "") for cell in raw):
            continue
        record = {headers[i]: raw[i] if i < len(raw) else None for i in range(len(headers)) if headers[i]}
        records.append(record)
    return records


def google_csv_url(kind: str) -> str | None:
    direct = os.environ.get(f"FNOMO_{kind}_CSV_URL")
    if direct:
        return direct
    direct = os.environ.get("FNOMO_GOOGLE_SHEET_CSV_URL") if kind == "MASTER_PIPELINE" else None
    if direct:
        return direct
    sheet_id = os.environ.get("FNOMO_GOOGLE_SHEET_ID")
    gid = os.environ.get(f"FNOMO_{kind}_GID")
    if sheet_id and gid:
        return f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?{urlencode({'format': 'csv', 'gid': gid})}"
    return None


def read_google_sheet_csv(url: str) -> list[dict[str, Any]]:
    with urlopen(url, timeout=30) as response:
        text = response.read().decode("utf-8-sig")
    reader = csv.DictReader(text.splitlines())
    return [dict(row) for row in reader if any((value or "").strip() for value in row.values())]


def canonical_leads(workbook: Path) -> list[dict[str, Any]]:
    csv_url = google_csv_url("MASTER_PIPELINE")
    source_rows = read_google_sheet_csv(csv_url) if csv_url else read_xlsx_sheet(workbook, "FNOMO_MASTER_PIPELINE", "Lead_ID")
    leads = []
    for row in source_rows:
        lead = {key: compact(value_by_alias(row, aliases)) for key, aliases in LEAD_FIELD_ALIASES.items()}
        if not lead["lead_id"]:
            continue
        lead["last_contact_date"] = parse_date(lead["last_contact_date"])
        lead["next_action_date"] = parse_date(lead["next_action_date"])
        lead["source"] = "Google Sheet" if csv_url else "Excel mirror"
        lead["sync_key"] = sync_key(lead["name"], lead["company"])
        lead["row_hash"] = row_hash(lead)
        leads.append(lead)
    return leads


def canonical_history(workbook: Path) -> list[dict[str, Any]]:
    csv_url = google_csv_url("EXECUTION_HISTORY")
    rows = read_google_sheet_csv(csv_url) if csv_url else read_xlsx_sheet(workbook, "FNOMO_EXECUTION_HISTORY", "Date")
    events = []
    for index, row in enumerate(rows, start=1):
        event = {key: compact(value_by_alias(row, aliases)) for key, aliases in HISTORY_FIELD_ALIASES.items()}
        event["date"] = parse_date(event["date"])
        if not any(event.values()):
            continue
        event["history_key"] = row_hash({"index": index, **event})
        event["row_hash"] = row_hash(event)
        events.append(event)
    return events


def row_hash(record: dict[str, Any]) -> str:
    encoded = json.dumps(record, sort_keys=True, ensure_ascii=True, default=str)
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()[:24]


def rich_text(value: Any) -> dict[str, Any]:
    text = compact(value)[:MAX_RICH_TEXT]
    return {"rich_text": [{"type": "text", "text": {"content": text}}]} if text else {"rich_text": []}


def title_text(value: Any) -> dict[str, Any]:
    text = compact(value)[:200] or "Untitled"
    return {"title": [{"type": "text", "text": {"content": text}}]}


def select_value(value: Any) -> dict[str, Any]:
    text = compact(value)[:100]
    return {"select": {"name": text}} if text else {"select": None}


def date_value(value: Any) -> dict[str, Any]:
    date = parse_date(value)
    return {"date": {"start": date}} if date else {"date": None}


def now_date_value() -> dict[str, Any]:
    return {"date": {"start": datetime.now(timezone.utc).isoformat()}}


def lead_properties(lead: dict[str, Any]) -> dict[str, Any]:
    return lead_properties_for_schema(lead, "Lead ID", {})


def set_property(props: dict[str, Any], schema: dict[str, Any], name: str, value: dict[str, Any]) -> None:
    if not schema or name in schema:
        props[name] = value


def status_value(value: Any) -> dict[str, Any]:
    text = compact(value)[:100]
    return {"status": {"name": text}} if text else {"status": None}


def text_property_for_schema(schema: dict[str, Any], name: str, value: Any) -> dict[str, Any]:
    if schema.get(name, {}).get("type") == "title":
        return title_text(value)
    return rich_text(value)


def lead_properties_for_schema(lead: dict[str, Any], title_property: str, schema: dict[str, Any]) -> dict[str, Any]:
    title_source = lead["company"] if title_property == "Company" else lead["lead_id"]
    props: dict[str, Any] = {
        title_property: title_text(title_source),
        "Lead ID": text_property_for_schema(schema, "Lead ID", lead["lead_id"]),
        "Sync Key": rich_text(lead["sync_key"]),
        "Name": rich_text(lead["name"]),
        "Type": select_value(lead["lead_type"]),
        "Tier": select_value(lead["tier"]),
        "Stage": select_value(lead["stage"]),
        "Last Contact Date": date_value(lead["last_contact_date"]),
        "Next Action": rich_text(lead["next_action"]),
        "Next Action Date": date_value(lead["next_action_date"]),
        "Status": status_value(stage_to_template_status(lead["stage"], lead["pipeline_status"])),
        "Pipeline Status": select_value(lead["pipeline_status"]),
        "Notes": rich_text(lead["notes"]),
        "Source": select_value(lead["source"]),
        "Row Hash": rich_text(lead["row_hash"]),
        "Last Synced At": now_date_value(),
    }
    if title_property != "Company":
        props["Company"] = rich_text(lead["company"])
    if title_property == "Company" and "Contact" in schema:
        props["Contact"] = rich_text(lead["name"])
    if "Priority" in schema:
        props["Priority"] = select_value(tier_to_priority(lead["tier"]))
    if "Last action" in schema:
        props["Last action"] = date_value(lead["last_contact_date"])
    if "Next action" in schema:
        props["Next action"] = date_value(lead["next_action_date"])
    return {name: value for name, value in props.items() if not schema or name in schema}


def tier_to_priority(tier: str) -> str:
    return {"A": "High", "B": "Medium", "C": "Low"}.get(compact(tier).upper(), "Medium")


def stage_to_template_status(stage: str, fallback: str) -> str:
    stage_name = compact(stage)
    if stage_name == "New":
        return "Lead"
    if stage_name in {"Contacted", "Engaged", "Qualified"}:
        return "Active"
    if stage_name == "Converted":
        return "Complete"
    if stage_name == "Lost":
        return "Churned"
    return compact(fallback) or "Lead"


def task_type(next_action: str) -> str:
    text = compact(next_action).casefold()
    if "call" in text:
        return "Call"
    if "follow" in text:
        return "Follow-Up"
    return "Outreach"


def task_key(lead: dict[str, Any]) -> str:
    return f"{lead['sync_key']}::next-action"


def task_properties_for_schema(lead: dict[str, Any], lead_page_id: str, title_property: str, schema: dict[str, Any]) -> dict[str, Any]:
    action = lead["next_action"] or f"Move {lead['name'] or lead['company']} to next stage"
    title = action[:160]
    props: dict[str, Any] = {
        title_property: title_text(title),
        "Sync Key": rich_text(task_key(lead)),
        "Type": select_value(task_type(action)),
        "Due date": date_value(lead["next_action_date"]),
        "Status": status_value("To Do"),
        "Priority": select_value(tier_to_priority(lead["tier"])),
        "Client": {"relation": [{"id": lead_page_id}]},
        "Lead Name": rich_text(lead["name"]),
        "Company": rich_text(lead["company"]),
        "Additional Information": rich_text(f"Stage: {lead['stage']} | Status: {lead['pipeline_status']} | Source: {lead['source']}"),
        "Last Synced At": now_date_value(),
    }
    return {name: value for name, value in props.items() if name in schema}


def history_properties(event: dict[str, Any]) -> dict[str, Any]:
    return {
        "History Key": title_text(event["history_key"]),
        "Date": date_value(event["date"]),
        "Action Type": rich_text(event["action_type"]),
        "Lead or Task": rich_text(event["lead_or_task"]),
        "Owner": rich_text(event["owner"]),
        "Channel": rich_text(event["channel"]),
        "Outcome": rich_text(event["outcome"]),
        "Next Move": rich_text(event["next_move"]),
        "Source": rich_text(event["source"]),
        "Row Hash": rich_text(event["row_hash"]),
        "Last Synced At": now_date_value(),
    }


def create_leads_database(parent_page_id: str) -> str:
    payload = {
        "parent": {"type": "page_id", "page_id": parent_page_id},
        "title": [{"type": "text", "text": {"content": "FNOMO_MASTER_PIPELINE"}}],
        "properties": {
            "Lead ID": {"title": {}},
            "Sync Key": {"rich_text": {}},
            "Name": {"rich_text": {}},
            "Company": {"rich_text": {}},
            "Type": {"select": {}},
            "Tier": {"select": {"options": [{"name": "A", "color": "red"}, {"name": "B", "color": "yellow"}, {"name": "C", "color": "gray"}]}},
            "Stage": {"select": {"options": [{"name": name, "color": color} for name, color in [("New", "gray"), ("Contacted", "blue"), ("Engaged", "green"), ("Qualified", "purple"), ("Converted", "green"), ("Lost", "red")]]}},
            "Last Contact Date": {"date": {}},
            "Next Action": {"rich_text": {}},
            "Next Action Date": {"date": {}},
            "Pipeline Status": {"select": {}},
            "Notes": {"rich_text": {}},
            "Source": {"select": {"options": [{"name": "Google Sheet", "color": "green"}, {"name": "Excel mirror", "color": "gray"}]}},
            "Row Hash": {"rich_text": {}},
            "Last Synced At": {"date": {}},
        },
    }
    return notion_request("POST", "databases", payload)["id"]


def create_history_database(parent_page_id: str) -> str:
    payload = {
        "parent": {"type": "page_id", "page_id": parent_page_id},
        "title": [{"type": "text", "text": {"content": "FNOMO_EXECUTION_HISTORY"}}],
        "properties": {
            "History Key": {"title": {}},
            "Date": {"date": {}},
            "Action Type": {"rich_text": {}},
            "Lead or Task": {"rich_text": {}},
            "Owner": {"rich_text": {}},
            "Channel": {"rich_text": {}},
            "Outcome": {"rich_text": {}},
            "Next Move": {"rich_text": {}},
            "Source": {"rich_text": {}},
            "Row Hash": {"rich_text": {}},
            "Last Synced At": {"date": {}},
        },
    }
    return notion_request("POST", "databases", payload)["id"]


def resolve_database_ids(create_missing: bool) -> tuple[str, str]:
    config = load_backend_config()
    leads_db = os.environ.get("NOTION_FNOMO_LEADS_DATABASE_ID") or config.get("leads_database_id")
    history_db = os.environ.get("NOTION_FNOMO_EXECUTION_HISTORY_DATABASE_ID") or config.get("execution_history_database_id")
    if leads_db and history_db:
        return leads_db, history_db
    if not create_missing:
        die(
            "Set NOTION_FNOMO_LEADS_DATABASE_ID and NOTION_FNOMO_EXECUTION_HISTORY_DATABASE_ID, "
            "or rerun with --create-databases and NOTION_FNOMO_PARENT_PAGE_ID."
        )
    parent = os.environ.get("NOTION_FNOMO_PARENT_PAGE_ID") or config.get("parent_page_id")
    if not parent:
        die("NOTION_FNOMO_PARENT_PAGE_ID is required when --create-databases is used.")
    if not leads_db:
        leads_db = create_leads_database(parent)
        print(f"Created leads database: {leads_db}")
    if not history_db:
        history_db = create_history_database(parent)
        print(f"Created execution history database: {history_db}")
    return leads_db, history_db


def title_property_name(database: dict[str, Any]) -> str:
    for name, prop in database.get("properties", {}).items():
        if prop.get("type") == "title":
            return name
    die("Notion database is missing a title property.")


def rename_database(database_id: str, title: str) -> None:
    notion_request("PATCH", f"databases/{database_id}", {"title": [{"type": "text", "text": {"content": title}}]})


def ensure_leads_database_schema(database_id: str) -> None:
    database = notion_request("GET", f"databases/{database_id}")
    properties = database.get("properties", {})
    missing: dict[str, Any] = {}
    for name, spec in {
        "Lead ID": {"rich_text": {}},
        "Name": {"rich_text": {}},
        "Type": {"select": {}},
        "Tier": {"select": {}},
        "Stage": {"select": {}},
        "Last Contact Date": {"date": {}},
        "Next Action": {"rich_text": {}},
        "Next Action Date": {"date": {}},
        "Pipeline Status": {"select": {}},
        "Channel": {"rich_text": {}},
        "Notes": {"rich_text": {}},
        "Row Hash": {"rich_text": {}},
        "Last Synced At": {"date": {}},
    }.items():
        if name not in properties:
            missing[name] = spec
    if "Sync Key" not in properties:
        missing["Sync Key"] = {"rich_text": {}}
    if missing:
        notion_request("PATCH", f"databases/{database_id}", {"properties": missing})


def ensure_tasks_database_schema(database_id: str) -> None:
    database = notion_request("GET", f"databases/{database_id}")
    properties = database.get("properties", {})
    missing: dict[str, Any] = {}
    for name, spec in {
        "Sync Key": {"rich_text": {}},
        "Lead Name": {"rich_text": {}},
        "Company": {"rich_text": {}},
        "Last Synced At": {"date": {}},
    }.items():
        if name not in properties:
            missing[name] = spec
    if missing:
        notion_request("PATCH", f"databases/{database_id}", {"properties": missing})


def query_all_pages(database_id: str, title_property: str) -> dict[str, str]:
    page_map: dict[str, str] = {}
    cursor = None
    while True:
        payload: dict[str, Any] = {"page_size": 100}
        if cursor:
            payload["start_cursor"] = cursor
        data = notion_request("POST", f"databases/{database_id}/query", payload)
        for page in data.get("results", []):
            prop = page.get("properties", {}).get(title_property, {}).get("title", [])
            if prop:
                key = "".join(part.get("plain_text", "") for part in prop).strip()
                if key:
                    page_map[key] = page["id"]
        if not data.get("has_more"):
            return page_map
        cursor = data.get("next_cursor")


def plain_property(page: dict[str, Any], name: str) -> str:
    prop = page.get("properties", {}).get(name, {})
    prop_type = prop.get("type")
    if prop_type in ("title", "rich_text"):
        return "".join(part.get("plain_text", "") for part in prop.get(prop_type, [])).strip()
    if prop_type == "select":
        return (prop.get("select") or {}).get("name", "").strip()
    if prop_type == "date":
        return ((prop.get("date") or {}).get("start") or "").strip()
    return ""


def query_lead_pages_by_name_company(database_id: str) -> dict[str, str]:
    page_map: dict[str, str] = {}
    cursor = None
    while True:
        payload: dict[str, Any] = {"page_size": 100}
        if cursor:
            payload["start_cursor"] = cursor
        data = notion_request("POST", f"databases/{database_id}/query", payload)
        for page in data.get("results", []):
            key = plain_property(page, "Sync Key") or sync_key(plain_property(page, "Name"), plain_property(page, "Company"))
            if key and key != "::" and key not in page_map:
                page_map[key] = page["id"]
        if not data.get("has_more"):
            return page_map
        cursor = data.get("next_cursor")


def query_pages_by_rich_text_key(database_id: str, key_property: str) -> dict[str, str]:
    page_map: dict[str, str] = {}
    cursor = None
    while True:
        payload: dict[str, Any] = {"page_size": 100}
        if cursor:
            payload["start_cursor"] = cursor
        data = notion_request("POST", f"databases/{database_id}/query", payload)
        for page in data.get("results", []):
            key = plain_property(page, key_property)
            if key and key not in page_map:
                page_map[key] = page["id"]
        if not data.get("has_more"):
            return page_map
        cursor = data.get("next_cursor")


def archive_unmatched_lead_pages(database_id: str, valid_keys: set[str]) -> int:
    archived = 0
    cursor = None
    while True:
        payload: dict[str, Any] = {"page_size": 100}
        if cursor:
            payload["start_cursor"] = cursor
        data = notion_request("POST", f"databases/{database_id}/query", payload)
        for page in data.get("results", []):
            key = plain_property(page, "Sync Key") or sync_key(plain_property(page, "Name") or plain_property(page, "Contact"), plain_property(page, "Company"))
            if key not in valid_keys:
                notion_request("PATCH", f"pages/{page['id']}", {"archived": True})
                archived += 1
        if not data.get("has_more"):
            return archived
        cursor = data.get("next_cursor")


def upsert_pages(database_id: str, existing: dict[str, str], title_key: str, rows: list[dict[str, Any]], prop_builder) -> tuple[int, int]:
    created = 0
    updated = 0
    for row in rows:
        key = row[title_key]
        properties = prop_builder(row)
        if key in existing:
            notion_request("PATCH", f"pages/{existing[key]}", {"properties": properties})
            updated += 1
        else:
            payload = {"parent": {"database_id": database_id}, "properties": properties}
            page = notion_request("POST", "pages", payload)
            existing[key] = page["id"]
            created += 1
    return created, updated


def sync_execution_tasks(tasks_db: str, leads_by_key: dict[str, str], leads: list[dict[str, Any]], archive_stale: bool) -> tuple[int, int]:
    ensure_tasks_database_schema(tasks_db)
    rename_database(tasks_db, "Execution")
    database = notion_request("GET", f"databases/{tasks_db}")
    schema = database.get("properties", {})
    title_name = title_property_name(database)
    existing_tasks = query_pages_by_rich_text_key(tasks_db, "Sync Key")
    task_rows = [lead for lead in leads if lead.get("next_action") and lead.get("next_action_date") and lead["sync_key"] in leads_by_key]
    created, updated = upsert_pages(
        tasks_db,
        existing_tasks,
        "task_sync_key",
        [{**lead, "task_sync_key": task_key(lead)} for lead in task_rows],
        lambda lead: task_properties_for_schema(lead, leads_by_key[lead["sync_key"]], title_name, schema),
    )
    if archive_stale:
        valid = {task_key(lead) for lead in task_rows}
        for key, page_id in list(existing_tasks.items()):
            if key not in valid:
                notion_request("PATCH", f"pages/{page_id}", {"archived": True})
    return created, updated


def main() -> int:
    parser = argparse.ArgumentParser(description="Sync Fnomo pipeline and execution history to Notion.")
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK), help="Fallback XLSX source when Google Sheet CSV env is not set.")
    parser.add_argument("--dry-run", action="store_true", help="Read and validate source rows without writing to Notion.")
    parser.add_argument("--create-databases", action="store_true", help="Create missing Notion databases under NOTION_FNOMO_PARENT_PAGE_ID.")
    parser.add_argument("--sync-execution", action="store_true", help="Sync next-action tasks into NOTION_FNOMO_EXECUTION_TASKS_DATABASE_ID.")
    parser.add_argument("--archive-unmatched-leads", action="store_true", help="Archive Notion lead rows not present in the sheet sync key set.")
    args = parser.parse_args()

    workbook = Path(args.workbook)
    leads = canonical_leads(workbook)
    history = canonical_history(workbook)
    print(f"Source leads: {len(leads)}")
    print(f"Source execution history events: {len(history)}")
    if leads:
        print(f"First lead: {leads[0]['lead_id']} | {leads[0]['name']} | {leads[0]['stage']} | next={leads[0]['next_action_date']}")

    if args.dry_run:
        print("Dry run only. No Notion writes performed.")
        return 0

    leads_db, history_db = resolve_database_ids(args.create_databases)
    config = load_backend_config()
    tasks_db = os.environ.get("NOTION_FNOMO_EXECUTION_TASKS_DATABASE_ID") or config.get("execution_tasks_database_id")
    rename_database(leads_db, "Leads")
    ensure_leads_database_schema(leads_db)
    leads_database = notion_request("GET", f"databases/{leads_db}")
    leads_schema = leads_database.get("properties", {})
    lead_title = title_property_name(leads_database)
    existing_leads = query_lead_pages_by_name_company(leads_db)
    existing_history = query_all_pages(history_db, "History Key")
    lead_created, lead_updated = upsert_pages(
        leads_db,
        existing_leads,
        "sync_key",
        leads,
        lambda lead: lead_properties_for_schema(lead, lead_title, leads_schema),
    )
    archived_leads = archive_unmatched_lead_pages(leads_db, {lead["sync_key"] for lead in leads}) if args.archive_unmatched_leads else 0
    existing_leads = query_lead_pages_by_name_company(leads_db)
    task_created = task_updated = 0
    if args.sync_execution:
        if not tasks_db:
            die("Set NOTION_FNOMO_EXECUTION_TASKS_DATABASE_ID or execution_tasks_database_id in data/fnomo_notion_backend.json.")
        task_created, task_updated = sync_execution_tasks(tasks_db, existing_leads, leads, archive_stale=True)
    history_created, history_updated = upsert_pages(history_db, existing_history, "history_key", history, history_properties)
    print(f"Leads created: {lead_created}, updated: {lead_updated}")
    if args.archive_unmatched_leads:
        print(f"Unmatched Notion lead rows archived: {archived_leads}")
    if args.sync_execution:
        print(f"Execution tasks created: {task_created}, updated: {task_updated}")
    print(f"History events created: {history_created}, updated: {history_updated}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
