#!/usr/bin/env python3
"""
Implement the Fnomo India Membership Founder Circle plan.

This does not restructure the CRM. It uses allowed operating fields:
- Tags
- Next Action / Next Action Date
- Notes
- Active Work tasks
- Execution History
- Founder Circle command page
"""

from __future__ import annotations

import json
import os
import sys
import time
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from shutil import copy2
from typing import Any

import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "EXCEL" / "War Room_ Community Outreach Pipeline - FNOMO Master.xlsx"
PARENT_PAGE_ID = "5818aa3c-b995-827e-af29-011d0732f5bc"
CLIENTS_DB = "e8a8aa3c-b995-8368-8495-014b8ea295c3"
CONTACTS_DB = "3548aa3c-b995-8148-a653-fead444e56fe"
EXECUTION_DB = "6fb8aa3c-b995-82ec-806b-01b6cd1b2915"
HISTORY_DB = "3548aa3c-b995-818a-8ede-ed9a5d99aac3"
NOTION_VERSION = "2022-06-28"
PAGE_TITLE = "Fnomo Founder Circle Operating Plan"

FOUNDER_TAGS = [
    "Messaging Test Batch 1",
    "Founder Circle Candidate",
    "Founder Circle Day 1-3",
    "Yearly Rs 4899",
    "No Cold Expansion",
]


SESSION = requests.Session()


def die(message: str) -> None:
    print(f"ERROR: {message}", file=sys.stderr)
    raise SystemExit(1)


def headers() -> dict[str, str]:
    token = os.environ.get("NOTION_API_KEY")
    if not token:
        die("NOTION_API_KEY missing")
    return {
        "Authorization": f"Bearer {token}",
        "Notion-Version": NOTION_VERSION,
        "Content-Type": "application/json",
    }


def notion(method: str, path: str, payload: dict[str, Any] | None = None) -> dict[str, Any]:
    for attempt in range(6):
        response = SESSION.request(
            method,
            f"https://api.notion.com/v1/{path.lstrip('/')}",
            headers=headers(),
            json=payload,
            timeout=60,
        )
        if response.status_code == 429:
            time.sleep(2 + attempt)
            continue
        if response.status_code >= 400:
            raise RuntimeError(f"{method} {path} failed {response.status_code}: {response.text[:1500]}")
        return response.json()
    raise RuntimeError(f"{method} {path} failed after retries")


def query_database(database_id: str, filter_payload: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    cursor: str | None = None
    while True:
        payload: dict[str, Any] = {"page_size": 100}
        if cursor:
            payload["start_cursor"] = cursor
        if filter_payload:
            payload["filter"] = filter_payload
        data = notion("POST", f"databases/{database_id}/query", payload)
        rows.extend(data.get("results", []))
        if not data.get("has_more"):
            return rows
        cursor = data.get("next_cursor")


def plain_text(prop: dict[str, Any] | None) -> str:
    if not prop:
        return ""
    prop_type = prop.get("type")
    if prop_type in {"title", "rich_text"}:
        return "".join(part.get("plain_text", "") for part in prop.get(prop_type, [])).strip()
    if prop_type == "select":
        return ((prop.get("select") or {}).get("name") or "").strip()
    if prop_type == "status":
        return ((prop.get("status") or {}).get("name") or "").strip()
    if prop_type == "multi_select":
        return "; ".join(item.get("name", "") for item in prop.get("multi_select", []))
    if prop_type == "date" and prop.get("date"):
        return (prop["date"].get("start") or "").strip()
    return ""


def prop(page: dict[str, Any], name: str) -> dict[str, Any] | None:
    return (page.get("properties") or {}).get(name)


def title_of(page: dict[str, Any] | None) -> str:
    if not page:
        return ""
    for value in (page.get("properties") or {}).values():
        if value.get("type") == "title":
            return plain_text(value)
    return ""


def page_name(page: dict[str, Any] | None) -> str:
    if not page:
        return ""
    return plain_text(prop(page, "Company")) or plain_text(prop(page, "Name")) or title_of(page)


def relation_ids(page: dict[str, Any], name: str) -> list[str]:
    return [item.get("id", "") for item in ((prop(page, name) or {}).get("relation") or []) if item.get("id")]


def rich_text(value: str) -> dict[str, Any]:
    return {"rich_text": [{"text": {"content": value[:1900]}}]} if value else {"rich_text": []}


def date_value(value: date) -> dict[str, Any]:
    return {"date": {"start": value.isoformat()}}


def merged_tags(existing_prop: dict[str, Any] | None, tags: list[str]) -> dict[str, Any]:
    existing: list[str] = []
    if existing_prop and existing_prop.get("type") == "multi_select":
        existing = [item["name"] for item in existing_prop.get("multi_select", []) if item.get("name")]
    for tag in tags:
        if tag not in existing:
            existing.append(tag)
    return {"multi_select": [{"name": tag} for tag in existing]}


def ensure_tag_options() -> None:
    colors = ["blue", "purple", "green", "orange"]
    for database_id in [CLIENTS_DB, CONTACTS_DB, EXECUTION_DB]:
        schema = notion("GET", f"databases/{database_id}").get("properties") or {}
        existing_options = []
        if schema.get("Tags", {}).get("type") == "multi_select":
            existing_options = [
                {"name": option["name"], "color": option.get("color", "default")}
                for option in schema["Tags"].get("multi_select", {}).get("options", [])
                if option.get("name")
            ]
        option_names = {option["name"] for option in existing_options}
        options = list(existing_options)
        for index, tag in enumerate(FOUNDER_TAGS):
            if tag not in option_names:
                options.append({"name": tag, "color": colors[index % len(colors)]})
        if schema.get("Tags", {}).get("type") != "multi_select":
            notion("PATCH", f"databases/{database_id}", {"properties": {"Tags": {"multi_select": {"options": options}}}})
        else:
            notion("PATCH", f"databases/{database_id}", {"properties": {"Tags": {"multi_select": {"options": options}}}})


def has_tag(page: dict[str, Any], tag: str) -> bool:
    return tag in plain_text(prop(page, "Tags"))


def founder_next_action(segment: str) -> str:
    if segment == "CA Firms & Associations":
        return (
            "Founder Circle invite: ask about a client investment decision that went wrong, then offer Rs 4,899/year founder access."
        )
    if segment == "Corporates & Business Owners":
        return (
            "Founder Circle invite: ask about the last capital decision that looked right but should have been stress-tested."
        )
    return "Founder Circle invite: test decision-validation trigger and ask for a simple yes/no reply."


def normalize(value: str) -> str:
    return " ".join((value or "").lower().split())


def workbook_founder_targets() -> set[tuple[str, str]]:
    if not WORKBOOK.exists():
        return set()
    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    ws = wb["FNOMO_MASTER_PIPELINE"]
    headers = {ws.cell(4, column).value: column for column in range(1, ws.max_column + 1) if ws.cell(4, column).value}
    if "Tags" not in headers:
        return set()
    targets: set[tuple[str, str]] = set()
    for row in range(5, ws.max_row + 1):
        tags = str(ws.cell(row, headers["Tags"]).value or "")
        if "Messaging Test Batch 1" not in tags:
            continue
        name = normalize(str(ws.cell(row, headers["Name"]).value or ""))
        company = normalize(str(ws.cell(row, headers["Company"]).value or ""))
        if name and company:
            targets.add((name, company))
    return targets


def apply_notion(today: date) -> dict[str, Any]:
    ensure_tag_options()
    workbook_targets = workbook_founder_targets()
    clients = query_database(CLIENTS_DB)
    contacts = query_database(CONTACTS_DB)
    tasks = query_database(EXECUTION_DB)
    client_by_id = {client["id"]: client for client in clients}
    task_by_contact: dict[str, dict[str, Any]] = {}
    for task in tasks:
        for contact_id in relation_ids(task, "Contact"):
            if has_tag(task, "Messaging Test Batch 1") or "Re-engage" in title_of(task):
                task_by_contact.setdefault(contact_id, task)

    counts: Counter = Counter()
    touched_clients: set[str] = set()
    segment_counts: Counter = Counter()
    for contact in contacts:
        client_ids = relation_ids(contact, "Company")
        if not client_ids or client_ids[0] not in client_by_id:
            continue
        client = client_by_id[client_ids[0]]
        target_key = (normalize(title_of(contact)), normalize(page_name(client)))
        if not has_tag(contact, "Messaging Test Batch 1") and target_key not in workbook_targets:
            continue
        segment = plain_text(prop(contact, "Segment")) or plain_text(prop(client, "Segment"))
        action = founder_next_action(segment)
        note = (
            f"[{today.isoformat()}] Founder Circle Candidate: Rs 4,899/year, 100-seat cap, "
            "reply generation before public scale; monthly later has no founder privileges."
        )
        notion(
            "PATCH",
            f"pages/{contact['id']}",
            {
                "properties": {
                    "Next Action": rich_text(action),
                    "Next Action Date": date_value(today),
                    "Notes": rich_text(append_note(plain_text(prop(contact, "Notes")), note)),
                    "Tags": merged_tags(prop(contact, "Tags"), FOUNDER_TAGS),
                }
            },
        )
        counts["contacts_tagged"] += 1
        segment_counts[segment or "Unsegmented"] += 1
        touched_clients.add(client["id"])

        task_note = (
            "Founder Circle Day 1-3. Send the founder-trigger message, classify the reply, "
            "and do not push a demo. Success is first reply signal. Price: Rs 4,899/year."
        )
        existing_task = task_by_contact.get(contact["id"])
        if existing_task:
            notion(
                "PATCH",
                f"pages/{existing_task['id']}",
                {
                    "properties": {
                        "Additional Information": rich_text(task_note),
                        "Due date": date_value(today),
                        "Tags": merged_tags(prop(existing_task, "Tags"), FOUNDER_TAGS),
                    }
                },
            )
            counts["tasks_updated"] += 1
        else:
            notion(
                "POST",
                "pages",
                {
                    "parent": {"database_id": EXECUTION_DB},
                    "properties": {
                        "Task": {"title": [{"text": {"content": f"Founder Circle invite with {title_of(contact)} ({page_name(client)})"}}]},
                        "Client": {"relation": [{"id": client["id"]}]},
                        "Contact": {"relation": [{"id": contact["id"]}]},
                        "Type": {"select": {"name": "Follow-Up"}},
                        "Priority": {"select": {"name": "High"}},
                        "Due date": date_value(today),
                        "Status": {"status": {"name": "To Do"}},
                        "Stage Context": {"select": {"name": "Contacted"}},
                        "Additional Information": rich_text(task_note),
                        "Tags": {"multi_select": [{"name": tag} for tag in FOUNDER_TAGS]},
                    },
                },
            )
            counts["tasks_created"] += 1

    for client_id in touched_clients:
        client = client_by_id[client_id]
        client_note = (
            f"[{today.isoformat()}] Founder Circle Candidate: included in first validation cohort; "
            "do not expand cold volume until reply rate clears 10%."
        )
        notion(
            "PATCH",
            f"pages/{client_id}",
            {
                "properties": {
                    "Next Action": rich_text("Founder Circle validation: secure first reply before yearly close attempt."),
                    "Next Action Date": date_value(today),
                    "Notes": rich_text(append_note(plain_text(prop(client, "Notes")), client_note)),
                    "Tags": merged_tags(prop(client, "Tags"), FOUNDER_TAGS),
                }
            },
        )
        counts["clients_tagged"] += 1

    create_or_replace_founder_page(today, counts, segment_counts)
    log_notion_history(today, counts, segment_counts)
    return {"counts": dict(counts), "segments": dict(segment_counts)}


def append_note(existing: str, marker: str) -> str:
    if marker in existing:
        return existing
    return (existing + " | " + marker).strip(" |")


def text(text_value: str) -> list[dict[str, Any]]:
    return [{"type": "text", "text": {"content": text_value[:1800]}}]


def block(block_type: str, text_value: str) -> dict[str, Any]:
    return {"object": "block", "type": block_type, block_type: {"rich_text": text(text_value)}}


def table(rows: list[list[Any]]) -> dict[str, Any]:
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    return {
        "object": "block",
        "type": "table",
        "table": {
            "table_width": width,
            "has_column_header": True,
            "has_row_header": False,
            "children": [
                {
                    "object": "block",
                    "type": "table_row",
                    "table_row": {"cells": [text(str(value)) for value in row]},
                }
                for row in normalized
            ],
        },
    }


def find_page(title: str) -> str | None:
    data = notion("POST", "search", {"query": title, "filter": {"value": "page", "property": "object"}, "page_size": 20})
    for result in data.get("results", []):
        if title_of(result).lower() == title.lower():
            return result["id"]
    return None


def archive_children(page_id: str) -> None:
    cursor: str | None = None
    while True:
        path = f"blocks/{page_id}/children?page_size=100"
        if cursor:
            path += f"&start_cursor={cursor}"
        data = notion("GET", path)
        for child in data.get("results", []):
            notion("PATCH", f"blocks/{child['id']}", {"archived": True})
        if not data.get("has_more"):
            break
        cursor = data.get("next_cursor")


def create_or_replace_founder_page(today: date, counts: Counter, segments: Counter) -> None:
    page_id = find_page(PAGE_TITLE)
    properties = {"title": {"title": [{"text": {"content": PAGE_TITLE}}]}}
    if page_id:
        notion("PATCH", f"pages/{page_id}", {"archived": False, "is_locked": False, "properties": properties})
        archive_children(page_id)
    else:
        page = notion("POST", "pages", {"parent": {"page_id": PARENT_PAGE_ID}, "properties": properties})
        page_id = page["id"]

    children = [
        block("paragraph", f"Founder Circle command layer generated {today.isoformat()}. CRM remains frozen; this page drives the Rs 4,899/year first cohort."),
        block("heading_1", "Positioning"),
        table(
            [
                ["Principle", "Operating Language"],
                ["Core frame", "Fnomo helps serious investors validate the thinking behind an investment before money moves."],
                ["Use", "decision validation, capital judgment, second layer of confidence, Founder Circle, first cohort"],
                ["Avoid", "tool, tips, demo, education, discount"],
            ]
        ),
        block("heading_1", "Founder Circle Rules"),
        table(
            [
                ["Rule", "Control"],
                ["Cohort cap", "100 Founder Circle members"],
                ["Founder price", "Rs 4,899/year; no discounting below this"],
                ["Public launch gate", "Do not open public sales until 20 paid founders"],
                ["Expansion gate", "No new cold expansion until reply rate exceeds 10% on current Tier A batch"],
                ["Proof gate", "Broader scale starts only after founder proof exists"],
            ]
        ),
        block("heading_1", "Founder Selection Criteria"),
        table(
            [
                ["Filter", "Strict Criteria"],
                ["Founder definition", "Not an early buyer. A founder has felt the cost of wrong decisions and is actively trying to reduce it."],
                ["Capital exposure", "Invests Rs 5L+ annually or influences capital decisions for clients, business, family, or network."],
                ["Decision pain", "Has made at least one wrong decision recently or uses many inputs but still lacks conviction."],
                ["Behavioral signal", "Engages with finance content, WhatsApp discussions, Telegram groups, or serious investor conversations."],
                ["Network leverage", "Can refer 2-5 similar people if the value is real."],
                ["Selection rule", "Do not invite everyone. Selection itself is the positioning."],
            ]
        ),
        block("heading_1", "Current Founder Validation Cohort"),
        table([["Segment", "Contacts"]] + [[segment, value] for segment, value in sorted(segments.items())]),
        block("heading_1", "Outreach Copy"),
        table(
            [
                ["Channel", "Copy"],
                [
                    "WhatsApp CA",
                    "Quick question - in the last few months, have you had a client take an investment decision that did not work out as expected? We are building Fnomo for exactly that gap: a validation layer before the client commits capital. We are opening the first Founder Circle for a small set of CAs. Would you be open to seeing the founder access?",
                ],
                [
                    "Business owner",
                    "What was the last capital or investment decision that looked right at the time, but you later wished you had stress-tested? Fnomo is built for that moment before money moves. We are inviting a small Founder Circle at Rs 4,899/year before public launch. Worth a quick look?",
                ],
                [
                    "LinkedIn",
                    "Most investment mistakes do not feel reckless when they are made. They feel informed. That is the gap Fnomo is building for: validating capital decisions before execution. We are opening a small Founder Circle for serious investors, advisors, and operators before public launch.",
                ],
            ]
        ),
        block("heading_1", "Daily Operating Targets"),
        table(
            [
                ["Metric", "Target"],
                ["Qualified founder conversations", "20-25/day"],
                ["Targeted messages", "70-90/day only after reply rate clears 10%"],
                ["Follow-ups", "40-60/day"],
                ["Referral asks", "10/day from warm or converted members"],
                ["Immediate success metric", "First 5-10 replies"],
            ]
        ),
        block("heading_1", "Social Proof Capture"),
        table(
            [
                ["Moment", "Ask"],
                ["Day 1-2", "What was your last wrong investment or capital decision?"],
                ["Day 3-5", "After using this, what changed in how you think before investing?"],
                ["Proof format", "1-line WhatsApp quote, 20-30 second voice note, LinkedIn founder story, or before/after case snippet."],
                ["Sales insertion", "One founder told me the biggest mistake was not wrong stock selection; it was not validating the decision before acting."],
                ["Rule", "Ask for decision clarity shift, not product praise."],
            ]
        ),
        block("heading_1", "Visibility And Momentum"),
        table(
            [
                ["Channel", "Operating Line"],
                ["LinkedIn daily", "Most investment mistakes do not feel reckless. They feel informed. That is the gap Fnomo is building for."],
                ["WhatsApp groups", "One thing I have noticed - people rarely validate decisions before investing. That is where most mistakes happen."],
                ["Follow-ups", "We are filling Founder Circle gradually and keeping it small intentionally for feedback quality."],
                ["Scarcity language", "Limited to 100 members; public access opens only after founder feedback stabilizes."],
                ["Avoid", "Countdown timers, fake urgency, last-chance spam, and inflated demand claims."],
            ]
        ),
        block("heading_1", "Launch Phases"),
        table(
            [
                ["Phase", "Goal", "Gate"],
                ["Days 1-5: Signal Validation", "5-10 strong conversations and 5+ replies with real pain", "If no replies, messaging is wrong."],
                ["Days 6-14: Founder Close Sprint", "20 paid founders from warm/referral/high-intent conversations", "Do not public launch before this."],
                ["Days 15-25: Proof Layer", "50 founders and 10-15 strong proof assets", "Use proof in LinkedIn and WhatsApp follow-ups."],
                ["Days 25-40: Scale Bridge", "3-5 sales/day", "Reply rate >15%, conversation conversion >25%, referrals emerging."],
            ]
        ),
        block("heading_1", "Objection Responses"),
        table(
            [
                ["Objection", "Response"],
                ["I'll try monthly first", "That is logical if this were a normal subscription. Founder Circle is different: early access, status, and locked founder privileges. If you are serious enough to test investment decisions, yearly is the cleaner choice."],
                ["What exactly do I get?", "A structured way to validate an investment decision before committing capital, plus founder-only feedback access and early workflows."],
                ["I don't have time", "That is exactly why this exists. It is for the moment before a capital decision, not for general learning."],
                ["Send details", "I'll send the founder note. Before I do, what kind of investment decision would you most want validated before acting?"],
            ]
        ),
        block("heading_1", "Advocate Loop"),
        table(
            [
                ["Step", "Script / Control"],
                ["Trigger usage", "What decision are you making next?"],
                ["Reinforce behavior", "Did you validate it before acting?"],
                ["Referral ask", "Who else do you know makes similar decisions?"],
                ["Recognition", "Founder referrals are status-based, not discount-based."],
                ["Daily proof target", "Capture 1-2 proof assets once real replies or value moments appear."],
            ]
        ),
        block("heading_1", "Next 24 Hours"),
        table(
            [
                ["Move", "Owner", "Stop Rule"],
                ["Send Founder Circle trigger to the 27 tagged Tier A contacts", "Kush", "Do not add cold leads before replies"],
                ["Classify every response in Outcome enum", "Kush/Codex", "No unclassified reply"],
                ["Ask every positive respondent for one decision they want validated", "Kush", "No demo-first flow"],
                ["Shortlist 20 next high-quality prospects only", "Kush", "Do not message them until reply-rate gate clears"],
                ["Refresh dashboard after replies", "Codex", "Old 0% metrics cannot remain after response data appears"],
            ]
        ),
    ]
    for index in range(0, len(children), 80):
        notion("PATCH", f"blocks/{page_id}/children", {"children": children[index : index + 80]})
    notion("PATCH", f"pages/{page_id}", {"is_locked": True})


def log_notion_history(today: date, counts: Counter, segments: Counter) -> None:
    key = f"{today.isoformat()}::founder-circle-plan-implemented"
    existing = query_database(HISTORY_DB, {"property": "History Key", "title": {"equals": key}})
    properties = {
        "History Key": {"title": [{"text": {"content": key}}]},
        "Date": date_value(today),
        "Action Type": rich_text("Founder Circle implementation"),
        "Lead or Task": rich_text("India Membership Rs 4,899/year Founder Circle"),
        "Owner": rich_text("Codex"),
        "Channel": rich_text("Workbook + Notion"),
        "Outcome": {"select": {"name": "No Response"}},
        "Outcome Notes": rich_text(
            f"Tagged founder validation cohort, updated active work, and created Founder Circle command page. Counts: {dict(counts)}; segments: {dict(segments)}"
        ),
        "Next Move": rich_text("Send Founder Circle triggers to the tagged cohort; success metric is first 5-10 replies."),
        "Source": rich_text("Founder Circle plan"),
        "Last Synced At": {"date": {"start": datetime.now().astimezone().isoformat(timespec="seconds")}},
    }
    if existing:
        notion("PATCH", f"pages/{existing[0]['id']}", {"properties": properties})
    else:
        notion("POST", "pages", {"parent": {"database_id": HISTORY_DB}, "properties": properties})


def apply_workbook(today: date) -> dict[str, Any]:
    backup = WORKBOOK.parent / "backups" / f"{WORKBOOK.stem}.before-founder-circle-plan.{datetime.now().strftime('%Y%m%d-%H%M%S')}{WORKBOOK.suffix}"
    backup.parent.mkdir(exist_ok=True)
    copy2(WORKBOOK, backup)
    wb = load_workbook(WORKBOOK)
    ws = wb["FNOMO_MASTER_PIPELINE"]
    header_row = 4
    headers = {ws.cell(header_row, column).value: column for column in range(1, ws.max_column + 1) if ws.cell(header_row, column).value}
    if "Tags" not in headers:
        headers["Tags"] = ws.max_column + 1
        ws.cell(header_row, headers["Tags"]).value = "Tags"

    updated = 0
    by_type: Counter = Counter()
    for row in range(header_row + 1, ws.max_row + 1):
        tags_cell = ws.cell(row, headers["Tags"])
        tags = str(tags_cell.value or "")
        if "Messaging Test Batch 1" not in tags:
            continue
        add_tags = ["Founder Circle Candidate", "Founder Circle Day 1-3", "Yearly Rs 4,899", "No Cold Expansion"]
        for tag in add_tags:
            if tag not in tags:
                tags = (tags + "; " + tag).strip("; ")
        tags_cell.value = tags
        lead_type = str(ws.cell(row, headers["Type (CA / Business / Association)"]).value or "").strip()
        if lead_type == "CA":
            action = "Founder Circle invite: ask whether any client made an investment decision that went wrong, then offer Rs 4,899/year founder access."
        elif lead_type == "Business":
            action = "Founder Circle invite: ask what capital decision looked right but should have been stress-tested before money moved."
        else:
            action = "Founder Circle invite: test decision-validation trigger and ask for a simple yes/no reply."
        ws.cell(row, headers["Next_Action"]).value = action
        ws.cell(row, headers["Next_Action_Date"]).value = today.isoformat()
        notes = str(ws.cell(row, headers["Notes (free text)"]).value or "")
        marker = f"[{today.isoformat()}] Founder Circle Candidate: Rs 4,899/year, 100-seat cap, first 5-10 replies before expansion."
        ws.cell(row, headers["Notes (free text)"]).value = append_note(notes, marker)
        updated += 1
        by_type[lead_type] += 1

    update_outreach_ready(wb, today)
    update_execution_tasks(wb, today)
    update_progress_sheet(wb, today, updated)
    append_workbook_history(wb, today, updated)
    wb.save(WORKBOOK)
    return {"backup": str(backup), "pipeline_rows_updated": updated, "by_type": dict(by_type)}


def update_outreach_ready(wb: Any, today: date) -> None:
    ws = wb["FNOMO_OUTREACH_READY"]
    header_row = 4
    headers = {ws.cell(header_row, column).value: column for column in range(1, ws.max_column + 1) if ws.cell(header_row, column).value}
    rows = {
        "Founder Circle - CA Invite": [
            "Founder Circle - CA Invite",
            "WhatsApp + Email",
            "Kush",
            today.isoformat(),
            "2026-05-05",
            "2026-05-07",
            "2026-05-10",
            "Founder Circle access for CAs",
            "Quick question - in the last few months, have you had a client take an investment decision that did not work out as expected?\n\nWe are building Fnomo for exactly that gap: a validation layer before the client commits capital.\n\nWe are opening the first Founder Circle for a small set of CAs at Rs 4,899/year before public launch.\n\nWould you be open to seeing the founder access?",
        ],
        "Founder Circle - Business Owner Invite": [
            "Founder Circle - Business Owner Invite",
            "WhatsApp + Email",
            "Kush",
            today.isoformat(),
            "2026-05-05",
            "2026-05-07",
            "2026-05-10",
            "What looked right before money moved?",
            "What was the last capital or investment decision that looked right at the time, but you later wished you had stress-tested?\n\nFnomo is built for that moment before money moves. We are inviting a small Founder Circle at Rs 4,899/year before public launch.\n\nWorth a quick look?",
        ],
        "Founder Circle - LinkedIn Post": [
            "Founder Circle - LinkedIn Post",
            "LinkedIn",
            "Kush",
            today.isoformat(),
            "",
            "",
            "",
            "Most investment mistakes do not feel reckless",
            "Most investment mistakes do not feel reckless when they are made. They feel informed.\n\nThat is the gap Fnomo is building for: validating capital decisions before execution.\n\nWe are opening a small Founder Circle for serious investors, advisors, and operators before public launch.",
        ],
        "Founder Circle - Objection Responses": [
            "Founder Circle - Objection Responses",
            "Sales Reply",
            "Kush",
            today.isoformat(),
            "",
            "",
            "",
            "Founder Circle objection handling",
            "Monthly first: That is logical if this were a normal subscription. Founder Circle is different: early access, status, and locked founder privileges.\n\nWhat exactly do I get? A structured way to validate an investment decision before committing capital, plus founder-only feedback access and early workflows.\n\nNo time: That is exactly why this exists. It is for the moment before a capital decision, not for general learning.",
        ],
        "Founder Circle - Proof Capture": [
            "Founder Circle - Proof Capture",
            "WhatsApp + Voice Note",
            "Kush",
            today.isoformat(),
            "",
            "",
            "",
            "Founder proof capture questions",
            "Day 1-2: What was your last wrong investment or capital decision?\n\nDay 3-5: After using this, what changed in how you think before investing?\n\nProof ask: Would you be comfortable sharing this as a 1-line WhatsApp quote or a 20-30 second voice note? We can keep it anonymized.",
        ],
        "Founder Circle - Referral Ask": [
            "Founder Circle - Referral Ask",
            "WhatsApp",
            "Kush",
            today.isoformat(),
            "",
            "",
            "",
            "Who else makes similar decisions?",
            "Who are two people you know who make serious investment or capital decisions and would value an extra validation layer before committing money?\n\nWe are keeping Founder Circle curated, so I would rather take 2 relevant introductions than broad sharing.",
        ],
    }
    existing = {str(ws.cell(row, headers["Lead_Name"]).value or ""): row for row in range(header_row + 1, ws.max_row + 1)}
    for lead_name, values in rows.items():
        row = existing.get(lead_name) or ws.max_row + 1
        for index, value in enumerate(values, 1):
            ws.cell(row, index).value = value
        existing[lead_name] = row


def update_execution_tasks(wb: Any, today: date) -> None:
    ws = wb["FNOMO_EXECUTION_TASKS"]
    header_row = 4
    headers = {ws.cell(header_row, column).value: column for column in range(1, ws.max_column + 1) if ws.cell(header_row, column).value}
    existing = {str(ws.cell(row, headers["Task"]).value or ""): row for row in range(header_row + 1, ws.max_row + 1)}
    tasks = [
        [
            "TDA-FC-001",
            "Send Founder Circle CA invite",
            "Kush",
            "A",
            "Send Founder Circle wrong-decision trigger to tagged CA contacts; classify every reply.",
            today.isoformat(),
            "Execute Today",
            "Founder Circle",
            "Target first 5-10 replies; do not open public sales.",
        ],
        [
            "TDA-FC-002",
            "Send Founder Circle business-owner invite",
            "Kush",
            "A",
            "Send capital-decision stress-test trigger to tagged business-owner contacts.",
            today.isoformat(),
            "Execute Today",
            "Founder Circle",
            "Use Rs 4,899/year founder access; no discount framing.",
        ],
        [
            "TDA-FC-003",
            "Publish Founder Circle LinkedIn post",
            "Kush",
            "B",
            "Publish serious-investor Founder Circle post after direct cohort messages are queued.",
            today.isoformat(),
            "Execute Today",
            "Founder Circle",
            "Do not imply fake demand; use real 100-seat cap.",
        ],
        [
            "TDA-FC-004",
            "Track Founder Circle replies",
            "Codex",
            "A",
            "Classify replies and refresh dashboard; success threshold is 5-10 replies.",
            today.isoformat(),
            "Execute Today",
            "Founder Circle",
            "No unclassified Outcome rows.",
        ],
        [
            "TDA-FC-005",
            "Shortlist 20 Founder Circle prospects",
            "Kush",
            "B",
            "Build a shortlist only: capital exposure, decision pain, behavioral signal, and referral leverage. Do not send until reply gate clears.",
            today.isoformat(),
            "Prepare Only",
            "Founder Circle",
            "This supports the daily stack without violating no cold expansion.",
        ],
        [
            "TDA-FC-006",
            "Capture Founder Circle proof assets",
            "Kush",
            "A",
            "After any positive reply/value moment, capture one line, voice note, or before/after thinking snippet.",
            today.isoformat(),
            "Execute Today",
            "Founder Circle",
            "Ask for decision clarity shift, not product praise.",
        ],
    ]
    for task in tasks:
        row = existing.get(task[1]) or ws.max_row + 1
        for index, value in enumerate(task, 1):
            ws.cell(row, index).value = value
        existing[task[1]] = row


def update_progress_sheet(wb: Any, today: date, founder_rows: int) -> None:
    if "FNOMO_CURRENT_PROGRESS" not in wb.sheetnames:
        return
    ws = wb["FNOMO_CURRENT_PROGRESS"]
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 1).value == "Founder Circle Control":
            ws.delete_rows(row, ws.max_row - row + 1)
            break
    start = ws.max_row + 3
    navy = "1F4E79"
    amber = "FFF2CC"
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def write(row: int, values: list[Any], fill: str | None = None, bold: bool = False) -> None:
        for column, value in enumerate(values, 1):
            cell = ws.cell(row, column)
            cell.value = value
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
            if bold:
                cell.font = Font(bold=True, color="FFFFFF" if fill == navy else "000000")

    write(start, ["Founder Circle Control", "Value", "Why It Matters", "Next Control"], navy, True)
    rows = [
        ["Founder Circle cap", "100 seats", "Real scarcity; not fake urgency", "Do not exceed cap"],
        ["Founder price", "Rs 4,899/year", "Yearly commitment is the conversion path", "No discount below founder price"],
        ["Founder validation cohort", founder_rows, "Existing Tier A no-response cohort", "Send before any cold expansion"],
        ["Founder selection filter", "Capital exposure + decision pain + network leverage", "Founder is a proof asset, not just an early buyer", "Reject weak-fit prospects"],
        ["Public launch gate", "20 paid founders", "Prevents public scaling before proof", "Hold broad market until reached"],
        ["Reply gate", "10% reply rate", "Confirms trigger works", "Current rate remains 0.0% until replies logged"],
        ["Proof target", "10-15 proof assets by Days 15-25", "Founder stories create the scale bridge", "Capture after value moments"],
    ]
    for index, row in enumerate(rows, start + 1):
        write(index, row, amber if row[0] in {"Reply gate", "Public launch gate"} else None)


def append_workbook_history(wb: Any, today: date, updated: int) -> None:
    ws = wb["FNOMO_EXECUTION_HISTORY"]
    row = ws.max_row + 1
    values = [
        today.isoformat(),
        "Founder Circle Implementation",
        "India Membership Rs 4,899/year",
        "Codex",
        "Workbook + Notion",
        f"Implemented Founder Circle operating layer; tagged {updated} existing Messaging Test contacts and added founder outreach/tasks.",
        "Send Founder Circle trigger; target first 5-10 replies before public expansion; shortlist only until reply gate clears.",
        "Founder Circle plan",
    ]
    for index, value in enumerate(values, 1):
        ws.cell(row, index).value = value


def main() -> None:
    today = date.fromisoformat(os.environ.get("FNOMO_OPERATING_DATE", date.today().isoformat()))
    notion_result = apply_notion(today)
    workbook_result = apply_workbook(today)
    print(json.dumps({"notion": notion_result, "workbook": workbook_result}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
