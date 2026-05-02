"""
war_room_updater.py
Reads and writes to the War Room Community Outreach Pipeline Excel file.
Appends completed leads to Outreach_Database sheet with correct column mapping.
"""

import os
import pandas as pd
from datetime import date
from pathlib import Path

WAR_ROOM_PATH = os.getenv(
    "WAR_ROOM_PATH",
    r"D:\Antigravity\eigent\Downloads\fnomo\EXCEL\War Room_ Community Outreach Pipeline.xlsx"
)
DB_SHEET = "Outreach_Database"


def _load_sheet(sheet_name: str) -> pd.DataFrame:
    try:
        return pd.read_excel(WAR_ROOM_PATH, sheet_name=sheet_name, engine="openpyxl")
    except Exception as e:
        print(f"[WAR ROOM] Could not load sheet '{sheet_name}': {e}")
        return pd.DataFrame()


def get_pending_corporate(n: int = 5) -> list[dict]:
    """Pull N pending corporate leads from Corporate Pilot Outreach sheet."""
    df = _load_sheet("Corporate Pilot Outreach")
    if df.empty:
        return []
    df.columns = [c.strip() for c in df.columns]
    status_col = next((c for c in df.columns if "status" in c.lower()), None)
    if status_col:
        df = df[df[status_col].astype(str).str.contains("New|new|pending|⚪", na=False)]
    return df.head(n).to_dict("records")


def append_results(results: list[dict]) -> bool:
    """
    Append completed leads into Outreach_Database sheet.
    Maps: firm name, email, status, focus area, outreach date, subject, scores.
    """
    try:
        from openpyxl import load_workbook
        wb = load_workbook(WAR_ROOM_PATH)

        if DB_SHEET not in wb.sheetnames:
            ws = wb.create_sheet(DB_SHEET)
            ws.append([
                "Name", "Email", "Group Name", "Focus Area",
                "Status", "Outreach Date", "Subject",
                "Hook Score", "Clarity Score", "Spam Risk", "Notes"
            ])
        else:
            ws = wb[DB_SHEET]

        today = date.today().strftime("%Y-%m-%d")
        for r in results:
            scores = r.get("scores", {})
            ws.append([
                r.get("name", r.get("lead_name", "")),
                r.get("email", r.get("lead_email", "")),
                r.get("company", r.get("name", "")),
                r.get("focus_area", "CA Partner Pipeline"),
                "Outreach Sent",
                today,
                r.get("subject", ""),
                scores.get("hook_strength", ""),
                scores.get("clarity_score", ""),
                scores.get("spam_risk", ""),
                r.get("status", ""),
            ])

        wb.save(WAR_ROOM_PATH)
        print(f"[WAR ROOM] Appended {len(results)} records to '{DB_SHEET}'")
        return True

    except Exception as e:
        print(f"[WAR ROOM] Write failed: {e}")
        return False
